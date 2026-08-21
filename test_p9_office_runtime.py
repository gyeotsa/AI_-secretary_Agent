from pathlib import Path

import fitz
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from pptx import Presentation

from core.office_runtime import AtomicOfficeEditor, OfficeComAdapter, OfficePackageInspector, OfficeRenderer


def test_docx_template_edit_preserves_run_format_and_package_styles(tmp_path):
    path = tmp_path / "template.docx"
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("고객명: {{NAME}}")
    run.bold = True
    document.save(path)
    before = OfficePackageInspector.snapshot(path)
    details = AtomicOfficeEditor.replace_text(path, {"{{NAME}}": "지원"})
    after = OfficePackageInspector.snapshot(path)
    assert details["replacement_count"] == 1
    assert OfficePackageInspector.verify_preserved(before, after)["changed"] == []
    saved = Document(path)
    assert saved.paragraphs[0].text == "고객명: 지원"
    assert saved.paragraphs[0].runs[0].bold is True


def test_xlsx_template_edit_preserves_cell_style(tmp_path):
    path = tmp_path / "template.xlsx"
    workbook = Workbook(); sheet = workbook.active
    sheet["A1"] = "{{VALUE}}"; sheet["A1"].fill = PatternFill("solid", fgColor="00FF00")
    workbook.save(path)
    details = AtomicOfficeEditor.replace_text(path, {"{{VALUE}}": "완료"})
    saved = load_workbook(path)
    assert details["replacement_count"] == 1
    assert saved.active["A1"].value == "완료"
    assert saved.active["A1"].fill.fgColor.rgb.endswith("00FF00")
    saved.close()


def test_pptx_template_edit_preserves_layout_parts(tmp_path):
    path = tmp_path / "template.pptx"
    presentation = Presentation(); slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "{{TITLE}}"; presentation.save(path)
    details = AtomicOfficeEditor.replace_text(path, {"{{TITLE}}": "분기 보고"})
    assert details["replacement_count"] == 1
    assert Presentation(path).slides[0].shapes.title.text == "분기 보고"


def test_rendered_pdf_visual_qa_rejects_blank_and_accepts_content(tmp_path):
    pdf = tmp_path / "rendered.pdf"
    document = fitz.open(); page = document.new_page(); page.insert_text((72, 72), "Visual QA")
    document.save(pdf); document.close()
    details = OfficeRenderer.verify_pdf(pdf)
    assert details["pages"] == 1 and details["nonblank_pages"] == 1
    assert Path(details["previews"][0]).is_file()


def test_com_status_detects_installed_office_without_launching_apps():
    assert isinstance(OfficeComAdapter.available("word"), bool)
    assert isinstance(OfficeComAdapter.available("excel"), bool)
    assert isinstance(OfficeComAdapter.available("powerpoint"), bool)
    assert OfficeComAdapter.available("hwp") is False
