"""Local, application-independent Excel workbook tools."""
from pathlib import Path
from typing import Any, Dict, List
import json

from openpyxl import Workbook, load_workbook
from core.harness import SafetyLayer
from core.plugin import BasePlugin, ToolSchema, IntentSchema, SlotSchema
from core.tool_result import Artifact, Evidence, ToolRunResult
from plugins._document_slots import extract_document_slots


class ExcelPlugin(BasePlugin):
    def __init__(self):
        super().__init__(); self.name = "excel"; self.description = "로컬 Excel 통합 문서 생성·조회·수정"

    def get_tools(self) -> List[ToolSchema]:
        return [
            ToolSchema("excel_create_workbook", "XLSX 통합 문서를 생성합니다", {"type":"object","properties":{"path":{"type":"string"},"sheet":{"type":"string","default":"Sheet1"},"rows":{"type":"array","items":{"type":"array"}}},"required":["path"]}, ["filesystem_write"]),
            ToolSchema("excel_read_workbook", "XLSX 시트 내용을 JSON으로 읽습니다", {"type":"object","properties":{"path":{"type":"string"},"sheet":{"type":"string"},"max_rows":{"type":"integer","default":200}},"required":["path"]}, ["filesystem_read"]),
            ToolSchema("excel_write_cells", "XLSX 셀 값을 수정합니다", {"type":"object","properties":{"path":{"type":"string"},"sheet":{"type":"string","default":"Sheet1"},"cells":{"type":"object"}},"required":["path","cells"]}, ["filesystem_write"]),
        ]
    def get_intents(self):
        return [IntentSchema("excel.create","Excel 통합 문서 생성","excel_create_workbook",["엑셀","xlsx","스프레드시트"],[SlotSchema("path","저장 경로","Excel 파일을 어디에 저장할까요, 보스?")],follow_up_hints=["다시"])]
    def extract_slots(self,intent_name,text,current_slots):
        return extract_document_slots(text,current_slots,".xlsx","새 통합 문서") if intent_name=="excel.create" else dict(current_slots)

    @staticmethod
    def _path(value: Any, suffix: str = ".xlsx") -> Path:
        path = Path(str(value)).expanduser().resolve()
        ok, error = SafetyLayer.validate_path(str(path))
        if not ok: raise ValueError(error)
        if path.suffix.casefold() != suffix: raise ValueError(f"{suffix} 파일만 지원합니다: {path}")
        return path

    def execute_tool(self, tool_name: str, data: Dict[str, Any]):
        try:
            path = self._path(data["path"])
            if tool_name == "excel_create_workbook":
                wb=Workbook(); ws=wb.active; ws.title=str(data.get("sheet") or "Sheet1")
                for row in data.get("rows") or []: ws.append(list(row))
                path.parent.mkdir(parents=True,exist_ok=True); wb.save(path)
                saved = load_workbook(path, read_only=True)
                try:
                    saved_sheet = saved[str(data.get("sheet") or "Sheet1")]
                    dimensions = {
                        "sheet": saved_sheet.title,
                        "rows": saved_sheet.max_row,
                        "columns": saved_sheet.max_column,
                    }
                finally:
                    saved.close()
                return ToolRunResult.successful(
                    tool_name=tool_name,
                    raw_output=f"Excel 파일 생성 성공: {path}",
                    evidence=[Evidence(
                        "xlsx_structure", "저장된 Excel 통합 문서와 시트를 다시 열어 확인했습니다.",
                        {"path": str(path), **dimensions, "size": path.stat().st_size},
                    )],
                    artifacts=[Artifact("spreadsheet", str(path), {"format": "xlsx"})],
                )
            wb=load_workbook(path)
            ws=wb[str(data.get("sheet"))] if data.get("sheet") else wb.active
            if tool_name == "excel_read_workbook":
                limit=max(1,min(int(data.get("max_rows",200)),5000)); rows=[list(r) for r in ws.iter_rows(max_row=limit,values_only=True)]
                return json.dumps({"path":str(path),"sheet":ws.title,"rows":rows},ensure_ascii=False,default=str)
            if tool_name == "excel_write_cells":
                for address,value in dict(data["cells"]).items(): ws[str(address)]=value
                wb.save(path)
                saved = load_workbook(path, read_only=True, data_only=False)
                try:
                    checked = saved[str(data.get("sheet") or "Sheet1")]
                    mismatches = {
                        address: checked[str(address)].value
                        for address, value in dict(data["cells"]).items()
                        if checked[str(address)].value != value
                    }
                finally:
                    saved.close()
                if mismatches:
                    raise ValueError(f"저장 후 셀 값 검증에 실패했습니다: {mismatches}")
                return ToolRunResult.successful(
                    tool_name=tool_name,
                    raw_output=f"Excel 셀 수정 성공: {path}",
                    evidence=[Evidence(
                        "xlsx_cells", "요청한 Excel 셀 값이 저장된 것을 확인했습니다.",
                        {"path": str(path), "cells": dict(data["cells"])},
                    )],
                    artifacts=[Artifact("spreadsheet", str(path), {"changed": True})],
                )
            return f"오류: 알 수 없는 툴 '{tool_name}'"
        except Exception as exc:
            if tool_name in {"excel_create_workbook", "excel_write_cells"}:
                return ToolRunResult.failed(tool_name=tool_name, error=str(exc))
            return f"오류: {exc}"
