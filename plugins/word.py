"""Local DOCX tools based on python-docx."""
from pathlib import Path
from typing import Any, Dict, List
from docx import Document
from core.harness import SafetyLayer
from core.plugin import BasePlugin, ToolSchema, IntentSchema, SlotSchema
from core.tool_result import Artifact, Evidence, ToolRunResult
from plugins._document_slots import extract_document_slots


class WordPlugin(BasePlugin):
    def __init__(self): super().__init__(); self.name="word"; self.description="로컬 Word 문서 생성·읽기"
    def get_tools(self) -> List[ToolSchema]:
        return [
            ToolSchema("word_create_document","DOCX 보고서를 생성합니다",{"type":"object","properties":{"path":{"type":"string"},"title":{"type":"string"},"paragraphs":{"type":"array","items":{"type":"string"}}},"required":["path"]},["filesystem_write"]),
            ToolSchema("word_read_document","DOCX 텍스트를 읽습니다",{"type":"object","properties":{"path":{"type":"string"}},"required":["path"]},["filesystem_read"]),
        ]
    def get_intents(self):
        return [IntentSchema("word.create","Word 문서 생성","word_create_document",["워드","word","docx"],[SlotSchema("path","저장 경로","Word 문서를 어디에 저장할까요, 보스?")],follow_up_hints=["다시"])]
    def extract_slots(self,intent_name,text,current_slots):
        return extract_document_slots(text,current_slots,".docx","새 문서") if intent_name=="word.create" else dict(current_slots)
    @staticmethod
    def _path(value: Any) -> Path:
        path=Path(str(value)).expanduser().resolve(); ok,error=SafetyLayer.validate_path(str(path))
        if not ok: raise ValueError(error)
        if path.suffix.casefold() != ".docx": raise ValueError(".docx 파일만 지원합니다")
        return path
    def execute_tool(self,name:str,data:Dict[str,Any]):
        try:
            path=self._path(data["path"])
            if name=="word_create_document":
                doc=Document()
                if data.get("title"): doc.add_heading(str(data["title"]),0)
                for text in data.get("paragraphs") or []: doc.add_paragraph(str(text))
                path.parent.mkdir(parents=True,exist_ok=True); doc.save(path)
                saved=Document(path)
                return ToolRunResult.successful(
                    tool_name=name,
                    raw_output=f"Word 문서 생성 성공: {path}",
                    evidence=[Evidence(
                        "docx_structure", "저장된 Word 문서를 다시 열어 문단 구조를 확인했습니다.",
                        {"path":str(path),"paragraphs":len(saved.paragraphs),"size":path.stat().st_size},
                    )],
                    artifacts=[Artifact("document",str(path),{"format":"docx"})],
                )
            if name=="word_read_document": return "\n".join(p.text for p in Document(path).paragraphs)
            return f"오류: 알 수 없는 툴 '{name}'"
        except Exception as exc:
            if name=="word_create_document": return ToolRunResult.failed(tool_name=name,error=str(exc))
            return f"오류: {exc}"
