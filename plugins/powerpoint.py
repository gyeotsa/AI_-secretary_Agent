"""Local PPTX tools based on python-pptx."""
from pathlib import Path
from typing import Any, Dict, List
import json
from pptx import Presentation
from core.harness import SafetyLayer
from core.plugin import BasePlugin, ToolSchema, IntentSchema, SlotSchema
from plugins._document_slots import extract_document_slots


class PowerPointPlugin(BasePlugin):
    def __init__(self): super().__init__(); self.name="powerpoint"; self.description="로컬 PowerPoint 생성·읽기"
    def get_tools(self)->List[ToolSchema]:
        return [
            ToolSchema("powerpoint_create_presentation","PPTX 발표 자료를 생성합니다",{"type":"object","properties":{"path":{"type":"string"},"title":{"type":"string"},"slides":{"type":"array","items":{"type":"object","properties":{"title":{"type":"string"},"bullets":{"type":"array","items":{"type":"string"}}}}}},"required":["path"]},["filesystem_write"]),
            ToolSchema("powerpoint_read_presentation","PPTX 텍스트를 읽습니다",{"type":"object","properties":{"path":{"type":"string"}},"required":["path"]},["filesystem_read"]),
        ]
    def get_intents(self):
        return [IntentSchema("powerpoint.create","PowerPoint 발표 자료 생성","powerpoint_create_presentation",["파워포인트","powerpoint","pptx","발표 자료"],[SlotSchema("path","저장 경로","PowerPoint 파일을 어디에 저장할까요, 보스?")],follow_up_hints=["다시"])]
    def extract_slots(self,intent_name,text,current_slots):
        return extract_document_slots(text,current_slots,".pptx","새 발표 자료") if intent_name=="powerpoint.create" else dict(current_slots)
    @staticmethod
    def _path(value:Any)->Path:
        path=Path(str(value)).expanduser().resolve(); ok,error=SafetyLayer.validate_path(str(path))
        if not ok: raise ValueError(error)
        if path.suffix.casefold() != ".pptx": raise ValueError(".pptx 파일만 지원합니다")
        return path
    def execute_tool(self,name:str,data:Dict[str,Any])->str:
        try:
            path=self._path(data["path"])
            if name=="powerpoint_create_presentation":
                prs=Presentation()
                if data.get("title"):
                    slide=prs.slides.add_slide(prs.slide_layouts[0]); slide.shapes.title.text=str(data["title"])
                for item in data.get("slides") or []:
                    slide=prs.slides.add_slide(prs.slide_layouts[1]); slide.shapes.title.text=str(item.get("title", "")); frame=slide.placeholders[1].text_frame; frame.clear()
                    for i,bullet in enumerate(item.get("bullets") or []): (frame.paragraphs[0] if i==0 else frame.add_paragraph()).text=str(bullet)
                path.parent.mkdir(parents=True,exist_ok=True); prs.save(path); return f"PowerPoint 생성 성공: {path}"
            if name=="powerpoint_read_presentation":
                slides=[]
                for slide in Presentation(path).slides: slides.append([shape.text for shape in slide.shapes if hasattr(shape,"text") and shape.text])
                return json.dumps(slides,ensure_ascii=False)
            return f"오류: 알 수 없는 툴 '{name}'"
        except Exception as exc:return f"오류: {exc}"
