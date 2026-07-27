import os
import subprocess
import tempfile
import json
import queue
import threading
import asyncio
import time
import hashlib
from typing import Optional
from datetime import datetime
from dataclasses import asdict
from pathlib import Path
from config import Config
from core.harness import SafetyLayer
from core.user_profile import get_user_profile
from core.workspace import get_workspace_manager
from core.plugin import get_plugin_registry
from core.tts_settings import get_tts_settings_manager
from core.custom_tts import GPTSoVITSClient, load_custom_voice_profiles, split_tts_text
from core.tts_normalizer import normalize_for_tts
from core.tool_result import Artifact, Evidence, ToolRunResult
from core.verifier import ToolVerifier

try:
    from ddgs import DDGS
except ImportError:
    try:
        from duckduckgo_search import DDGS
    except ImportError:
        DDGS = None

try:
    import pyttsx3
except ImportError:
    pyttsx3 = None

try:
    import edge_tts
except ImportError:
    edge_tts = None

try:
    import whisper
    import sounddevice as sd
    import numpy as np
    import scipy.io.wavfile as wav

    WHISPER_AVAILABLE = True
except ImportError:
    WHISPER_AVAILABLE = False


class ToolExecutor:
    def __init__(self):
        self.safety = SafetyLayer()
        self.user_profile = get_user_profile()
        self.workspace = get_workspace_manager()
        self.plugin_registry = get_plugin_registry()
        self.verifier = ToolVerifier()

        # Lazy initialization for optional modules
        self._rag_manager = None
        self._scheduler_manager = None
        self._hardware_manager = None
        self._multimodal_manager = None
        self._project_indexer = None

        # TTS engine
        self._tts_engine = None
        self._tts_lock = threading.Lock()
        self.tts_settings = get_tts_settings_manager()
        self._custom_tts_clients = {}

        # Load plugins from plugins directory
        try:
            self.plugin_registry.load_plugins_from_directory()
        except Exception as e:
            print(f"[ToolExecutor] Plugin 로딩 오류: {e}")

    @property
    def rag_manager(self):
        if self._rag_manager is None:
            try:
                from core.rag import get_rag_manager
                self._rag_manager = get_rag_manager()
            except Exception as e:
                self._rag_manager = None
        return self._rag_manager

    @property
    def scheduler_manager(self):
        if self._scheduler_manager is None:
            try:
                from core.scheduler import get_scheduler_manager
                self._scheduler_manager = get_scheduler_manager()
            except Exception as e:
                self._scheduler_manager = None
        return self._scheduler_manager

    @property
    def hardware_manager(self):
        if self._hardware_manager is None:
            try:
                from core.hardware import get_hardware_manager
                self._hardware_manager = get_hardware_manager()
            except Exception as e:
                self._hardware_manager = None
        return self._hardware_manager

    @property
    def multimodal_manager(self):
        if self._multimodal_manager is None:
            try:
                from core.multimodal import get_multimodal_manager
                self._multimodal_manager = get_multimodal_manager()
            except Exception as e:
                self._multimodal_manager = None
        return self._multimodal_manager

    @property
    def project_indexer(self):
        if self._project_indexer is None:
            try:
                from core.project_indexer import get_project_indexer
                self._project_indexer = get_project_indexer()
            except Exception as e:
                self._project_indexer = None
        return self._project_indexer

    def _resolve_and_validate_path(self, path: str) -> tuple[bool, str, str]:
        """경로를 Workspace나 SafetyLayer로 검증하고 유효한 절대 경로 반환"""
        resolved_path = path
        try:
            if self.workspace.is_set():
                real_path = self.workspace.resolve(path)
                resolved_path = str(real_path)
        except Exception as e:
            return False, f"오류: {e}", ""
        
        is_valid, error_msg = self.safety.validate_path(resolved_path)
        if not is_valid:
            return False, f"오류: {error_msg}", ""
        
        return True, "", resolved_path

    def read_file(self, path: str):
        is_valid, error_msg, resolved_path = self._resolve_and_validate_path(path)
        if not is_valid:
            return ToolRunResult.failed(tool_name="read_file",error=error_msg,raw_output=error_msg)

        try:
            content = Path(resolved_path).read_text(encoding="utf-8")
            return ToolRunResult.successful(
                tool_name="read_file",raw_output=content,
                evidence=[Evidence("file_content","파일 내용·크기·해시를 확인했습니다.",{
                    "path":resolved_path,"size":Path(resolved_path).stat().st_size,
                    "sha256":hashlib.sha256(content.encode("utf-8")).hexdigest(),
                })],
                artifacts=[Artifact("file",resolved_path)],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="read_file",error=f"파일 읽기 오류: {e}")

    def write_file(self, path: str, content: str):
        is_valid, error_msg, resolved_path = self._resolve_and_validate_path(path)
        if not is_valid:
            return ToolRunResult.failed(tool_name="write_file", error=error_msg)

        try:
            target = Path(resolved_path)
            before_hash = (
                hashlib.sha256(target.read_bytes()).hexdigest()
                if target.is_file()
                else None
            )
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(content, encoding="utf-8")
            saved_bytes = target.read_bytes()
            expected_bytes = content.encode("utf-8")
            if saved_bytes != expected_bytes:
                return ToolRunResult.failed(
                    tool_name="write_file",
                    error="파일 저장 후 내용 검증에 실패했습니다.",
                    evidence=[Evidence("file_content_mismatch", "저장된 바이트가 요청 내용과 다릅니다.", {
                        "path": str(target),
                        "expected_size": len(expected_bytes),
                        "actual_size": len(saved_bytes),
                    })],
                )
            after_hash = hashlib.sha256(saved_bytes).hexdigest()
            return ToolRunResult.successful(
                tool_name="write_file",
                raw_output=f"파일이 성공적으로 저장되었습니다: {target}",
                evidence=[Evidence("file_content", "저장 후 파일 내용과 해시를 확인했습니다.", {
                    "path": str(target),
                    "size": len(saved_bytes),
                    "before_sha256": before_hash,
                    "after_sha256": after_hash,
                    "changed": before_hash != after_hash,
                })],
                artifacts=[Artifact("file", str(target), {"sha256": after_hash})],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="write_file", error=f"파일 쓰기 오류: {e}")

    def list_directory(self, path: str = ""):
        target_path = path
        try:
            if not path and self.workspace.is_set():
                real_path = self.workspace.get_workspace_path()
                if real_path:
                    target_path = real_path
            elif self.workspace.is_set():
                real_path = self.workspace.resolve(path)
                target_path = str(real_path)
        except Exception as e:
            return ToolRunResult.failed(tool_name="list_directory",error=str(e))

        is_valid, error_msg = self.safety.validate_path(target_path)
        if not is_valid:
            return ToolRunResult.failed(tool_name="list_directory",error=error_msg)

        try:
            items = os.listdir(target_path)
            result = []
            for item in items:
                item_path = os.path.join(target_path, item)
                is_dir = os.path.isdir(item_path)
                prefix = "[폴더] " if is_dir else "[파일] "
                result.append(prefix + item)
            output = "\n".join(result)
            return ToolRunResult.successful(
                tool_name="list_directory",raw_output=output,
                evidence=[Evidence("directory_listing",f"디렉터리 항목 {len(items)}개를 조회했습니다.",{
                    "path":str(Path(target_path).resolve()),"count":len(items),
                })],
                artifacts=[Artifact("directory",str(Path(target_path).resolve()))],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="list_directory",error=f"디렉토리 목록 오류: {e}")

    def create_directory(self, dir_path: str):
        is_valid, error_msg, resolved_path = self._resolve_and_validate_path(dir_path)
        if not is_valid:
            return ToolRunResult.failed(tool_name="create_directory", error=error_msg)

        try:
            target = Path(resolved_path)
            existed = target.is_dir()
            target.mkdir(parents=True, exist_ok=True)
            if not target.is_dir():
                return ToolRunResult.failed(
                    tool_name="create_directory",
                    error="폴더 생성 후 존재 여부를 확인하지 못했습니다.",
                )
            return ToolRunResult.successful(
                tool_name="create_directory",
                raw_output=f"폴더가 성공적으로 생성되었습니다: {target}",
                evidence=[Evidence("directory_exists", "생성 후 폴더 존재를 확인했습니다.", {
                    "path": str(target),
                    "already_existed": existed,
                })],
                artifacts=[Artifact("directory", str(target))],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="create_directory", error=f"폴더 생성 오류: {e}")

    def delete_directory(self, dir_path: str):
        is_valid, error_msg, resolved_path = self._resolve_and_validate_path(dir_path)
        if not is_valid:
            return ToolRunResult.failed(tool_name="delete_directory", error=error_msg)

        try:
            target = Path(resolved_path)
            if not target.is_dir():
                return ToolRunResult.failed(
                    tool_name="delete_directory",
                    error=f"폴더를 찾을 수 없습니다: {target}",
                )

            import shutil
            shutil.rmtree(target)
            if target.exists():
                return ToolRunResult.failed(
                    tool_name="delete_directory",
                    error="폴더 삭제 후에도 대상 경로가 남아 있습니다.",
                )
            return ToolRunResult.successful(
                tool_name="delete_directory",
                raw_output=f"폴더가 성공적으로 삭제되었습니다: {target}",
                evidence=[Evidence("directory_absent", "삭제 후 대상 경로가 존재하지 않음을 확인했습니다.", {
                    "path": str(target),
                })],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="delete_directory", error=f"폴더 삭제 오류: {e}")

    # ------------------------------
    # Workspace 관련 도구 추가
    # ------------------------------
    def set_workspace(self, path: str):
        """작업 공간을 설정합니다."""
        try:
            success = self.workspace.set_workspace(path)
            if success:
                info = self.workspace.get_info()
                return ToolRunResult.successful(
                    tool_name="set_workspace",
                    raw_output=f"Workspace가 설정되었습니다: {info.name} (파일: {info.file_count}개)",
                    evidence=[Evidence("workspace_state","Workspace 설정과 실제 디렉터리를 확인했습니다.",{
                        "path":info.path,"name":info.name,"file_count":info.file_count,
                    })],
                    artifacts=[Artifact("directory",info.path,{"workspace":True})],
                )
            else:
                return ToolRunResult.failed(tool_name="set_workspace",error=f"유효하지 않은 경로입니다: {path}")
        except Exception as e:
            return ToolRunResult.failed(tool_name="set_workspace",error=f"Workspace 설정 오류: {e}")

    def get_workspace_info(self):
        """현재 Workspace 정보를 반환합니다."""
        if self.workspace.is_set():
            info = self.workspace.get_info()
            output = (f"Workspace: {info.name}\n"
                    f"경로: {info.path}\n"
                    f"파일 개수: {info.file_count}")
            return ToolRunResult.successful(
                tool_name="get_workspace_info",raw_output=output,
                evidence=[Evidence("workspace_state","현재 Workspace 정보를 확인했습니다.",asdict(info))],
                artifacts=[Artifact("directory",info.path,{"workspace":True})],
            )
        else:
            return ToolRunResult.failed(tool_name="get_workspace_info",error="Workspace가 설정되지 않았습니다.")

    def get_workspace_tree(self):
        """Workspace의 파일 트리를 반환합니다."""
        if not self.workspace.is_set():
            return ToolRunResult.failed(tool_name="get_workspace_tree",error="Workspace가 설정되지 않았습니다.")

        try:
            import json
            tree = self.workspace.get_file_tree()
            output = json.dumps(tree, indent=2, ensure_ascii=False)
            path = str(self.workspace.get_workspace_path())
            return ToolRunResult.successful(
                tool_name="get_workspace_tree",raw_output=output,
                evidence=[Evidence("workspace_tree","Workspace 파일 트리 생성을 확인했습니다.",{"path":path,"root":tree.get("name")})],
                artifacts=[Artifact("directory",path,{"workspace":True})],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="get_workspace_tree",error=f"트리 가져오기 오류: {e}")

    # ------------------------------
    # Project Indexer 관련 도구 추가
    # ------------------------------
    def index_project(self):
        """현재 Workspace나 프로젝트 루트를 인덱싱합니다."""
        if self.project_indexer is None:
            return ToolRunResult.failed(tool_name="index_project",error="Project Indexer를 초기화할 수 없습니다.")

        try:
            if not self.project_indexer.project_root and self.workspace.is_set():
                self.project_indexer.set_project_root(self.workspace.current_workspace)

            if not self.project_indexer.project_root:
                return ToolRunResult.failed(tool_name="index_project",error="프로젝트 루트가 설정되지 않았습니다. 먼저 Workspace를 설정하세요.")

            count = self.project_indexer.index_project()
            return ToolRunResult.successful(
                tool_name="index_project",raw_output=f"프로젝트 인덱싱 완료: {count}개 파일",
                evidence=[Evidence("project_index","프로젝트 인덱스 파일 수를 확인했습니다.",{"project_root":self.project_indexer.project_root,"file_count":count})],
                artifacts=[Artifact("project_index",self.project_indexer.db_path,{"project_root":self.project_indexer.project_root})],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="index_project",error=f"프로젝트 인덱싱 오류: {e}")

    def search_files(self, query: str, search_type: str = "name"):
        """파일을 검색합니다 (이름, 내용, 확장자)."""
        if self.project_indexer is None:
            return ToolRunResult.failed(tool_name="search_files",error="Project Indexer를 초기화할 수 없습니다.")

        try:
            results = self.project_indexer.search_files(query, search_type)
            output = json.dumps(results, indent=2, ensure_ascii=False) if results else "검색 결과가 없습니다."
            return ToolRunResult.successful(
                tool_name="search_files",raw_output=output,
                evidence=[Evidence("project_index_query",f"인덱스에서 파일 {len(results)}건을 조회했습니다.",{
                    "query":query,"search_type":search_type,"count":len(results),
                    "project_root":self.project_indexer.project_root,
                })],
                artifacts=[Artifact("file",str(item["path"])) for item in results if item.get("path")],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="search_files",error=f"파일 검색 오류: {e}")

    def search_symbols(self, query: str, symbol_type: Optional[str] = None):
        """코드 심볼을 검색합니다 (함수, 클래스, 변수)."""
        if self.project_indexer is None:
            return ToolRunResult.failed(tool_name="search_symbols",error="Project Indexer를 초기화할 수 없습니다.")

        try:
            results = self.project_indexer.search_symbols(query, symbol_type)
            output = json.dumps(results, indent=2, ensure_ascii=False) if results else "검색 결과가 없습니다."
            return ToolRunResult.successful(
                tool_name="search_symbols",raw_output=output,
                evidence=[Evidence("project_symbol_query",f"인덱스에서 코드 심볼 {len(results)}건을 조회했습니다.",{
                    "query":query,"symbol_type":symbol_type,"count":len(results),
                    "project_root":self.project_indexer.project_root,
                })],
                artifacts=[Artifact("file",str(item["file_path"])) for item in results if item.get("file_path")],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="search_symbols",error=f"심볼 검색 오류: {e}")

    def get_file_info(self, path: str):
        """파일의 상세 정보를 반환합니다."""
        if self.project_indexer is None:
            return ToolRunResult.failed(tool_name="get_file_info",error="Project Indexer를 초기화할 수 없습니다.")

        try:
            file_info = self.project_indexer.get_file_info(path)
            if not file_info:
                return ToolRunResult.failed(tool_name="get_file_info",error=f"파일 정보를 찾을 수 없습니다: {path}")

            import json
            output = json.dumps({
                "path": file_info.path,
                "name": file_info.name,
                "extension": file_info.extension,
                "size": file_info.size,
                "modified_at": datetime.fromtimestamp(
                    file_info.modified_at).isoformat() if file_info.modified_at else "",
                "is_text": file_info.is_text,
                "content_preview": file_info.content_preview,
                "symbols": file_info.symbols
            }, indent=2, ensure_ascii=False)
            return ToolRunResult.successful(
                tool_name="get_file_info",raw_output=output,
                evidence=[Evidence("project_file_info","인덱스 파일 메타데이터를 확인했습니다.",{
                    "path":file_info.path,"size":file_info.size,"modified_at":file_info.modified_at,
                })],
                artifacts=[Artifact("file",file_info.path)],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="get_file_info",error=f"파일 정보 가져오기 오류: {e}")

    def get_project_tree(self):
        """프로젝트의 파일 트리 구조를 반환합니다."""
        if self.project_indexer is None:
            return ToolRunResult.failed(tool_name="get_project_tree",error="Project Indexer를 초기화할 수 없습니다.")

        try:
            tree = self.project_indexer.get_file_tree()
            import json
            output = json.dumps(tree, indent=2, ensure_ascii=False)
            return ToolRunResult.successful(
                tool_name="get_project_tree",raw_output=output,
                evidence=[Evidence("project_tree","프로젝트 루트에서 파일 트리를 생성했습니다.",{"project_root":self.project_indexer.project_root})],
                artifacts=[Artifact("directory",str(self.project_indexer.project_root))],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="get_project_tree",error=f"프로젝트 트리 가져오기 오류: {e}")

    # ------------------------------
    # Memory 관련 도구 추가
    # ------------------------------
    def add_semantic_memory(self, key: str, content: str, category: str = "기타"):
        """시맨틱 메모리 추가"""
        try:
            from core.memory import get_semantic_memory, SemanticMemory
            semantic_manager = get_semantic_memory()
            memory = SemanticMemory(
                key=key,
                content=content,
                category=category
            )
            semantic_manager.add_memory(memory)
            saved = semantic_manager.get_memory(key)
            if not saved or saved.content != content or saved.category != category:
                raise ValueError("시맨틱 메모리 저장 후 재조회 검증에 실패했습니다.")
            return ToolRunResult.successful(
                tool_name="add_semantic_memory",raw_output=f"시맨틱 메모리가 추가되었습니다: {key}",
                evidence=[Evidence("semantic_memory","시맨틱 메모리 저장 후 키·카테고리·내용을 재확인했습니다.",{"key":key,"category":category,"content_sha256":hashlib.sha256(content.encode("utf-8")).hexdigest()})],
                artifacts=[Artifact("semantic_memory",key,{"category":category})],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="add_semantic_memory",error=f"시맨틱 메모리 추가 오류: {e}")

    def get_semantic_memory(self, key: str):
        """시맨틱 메모리 조회"""
        try:
            from core.memory import get_semantic_memory
            semantic_manager = get_semantic_memory()
            memory = semantic_manager.get_memory(key)
            if memory:
                return ToolRunResult.successful(
                    tool_name="get_semantic_memory",raw_output=f"{memory.key} ({memory.category}): {memory.content}",
                    evidence=[Evidence("semantic_memory","키로 시맨틱 메모리를 조회했습니다.",{"key":memory.key,"category":memory.category,"timestamp":memory.timestamp})],
                    artifacts=[Artifact("semantic_memory",memory.key,{"category":memory.category})],
                )
            else:
                return ToolRunResult.failed(tool_name="get_semantic_memory",error=f"메모리를 찾을 수 없습니다: {key}")
        except Exception as e:
            return ToolRunResult.failed(tool_name="get_semantic_memory",error=f"시맨틱 메모리 조회 오류: {e}")

    def search_semantic_memory(self, query: str, category: Optional[str] = None):
        """시맨틱 메모리 검색"""
        try:
            from core.memory import get_semantic_memory
            semantic_manager = get_semantic_memory()
            memories = semantic_manager.search_memories(query, category=category)
            result = []
            for mem in memories:
                result.append(f"- [{mem.category}] {mem.key}: {mem.content}")
            output = "\n".join(result) if result else "검색 결과가 없습니다."
            return ToolRunResult.successful(
                tool_name="search_semantic_memory",raw_output=output,
                evidence=[Evidence("semantic_memory_query",f"시맨틱 메모리 {len(memories)}건을 조회했습니다.",{"query":query,"category":category,"count":len(memories)})],
                artifacts=[Artifact("semantic_memory",mem.key,{"category":mem.category}) for mem in memories],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="search_semantic_memory",error=f"시맨틱 메모리 검색 오류: {e}")

    def delete_semantic_memory(self, key: str):
        """시맨틱 메모리 삭제"""
        try:
            from core.memory import get_semantic_memory
            semantic_manager = get_semantic_memory()
            deleted = semantic_manager.delete_memory(key)
            if deleted and semantic_manager.get_memory(key) is None:
                return ToolRunResult.successful(
                    tool_name="delete_semantic_memory",raw_output=f"메모리가 삭제되었습니다: {key}",
                    evidence=[Evidence("semantic_memory_absent","삭제 후 시맨틱 메모리가 존재하지 않음을 확인했습니다.",{"key":key})],
                )
            return ToolRunResult.failed(tool_name="delete_semantic_memory",error=f"메모리를 찾을 수 없습니다: {key}")
        except Exception as e:
            return ToolRunResult.failed(tool_name="delete_semantic_memory",error=f"시맨틱 메모리 삭제 오류: {e}")

    # ------------------------------
    # Automation Engine 관련 도구 추가
    # ------------------------------
    def add_automation_job(self, description: str, schedule_type: str, schedule_value: str, prompt: str):
        """자동화 작업 추가"""
        try:
            from core.scheduler import get_automation_engine
            engine = get_automation_engine()
            before_ids = {item["id"] for item in engine.get_job_records()}
            raw = engine.add_job(description, schedule_type, schedule_value, prompt)
            created = [
                item for item in engine.get_job_records()
                if item["id"] not in before_ids
                and item["description"] == description
                and item["schedule_type"] == schedule_type
                and str(item["schedule_value"]) == str(schedule_value)
                and item["prompt"] == prompt
            ]
            if not created:
                return ToolRunResult.failed(
                    tool_name="add_automation_job",
                    error=str(raw),
                    raw_output=str(raw),
                    evidence=[Evidence("scheduler_database", "등록 후 일치하는 작업 레코드를 찾지 못했습니다.")],
                )
            record = created[0]
            return ToolRunResult.successful(
                tool_name="add_automation_job",
                raw_output=str(raw),
                evidence=[Evidence("scheduler_database", "등록된 자동화 작업을 DB에서 다시 조회했습니다.", record)],
                artifacts=[Artifact("automation_job", str(record["id"]))],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="add_automation_job", error=f"자동화 작업 추가 오류: {e}")

    def list_automation_jobs(self):
        """자동화 작업 목록 보기"""
        try:
            from core.scheduler import get_automation_engine
            engine = get_automation_engine()
            records = engine.get_job_records()
            return ToolRunResult.successful(
                tool_name="list_automation_jobs",
                raw_output=engine.list_jobs(),
                evidence=[Evidence("scheduler_database", f"자동화 작업 {len(records)}건을 DB에서 조회했습니다.", {
                    "count": len(records),
                    "job_ids": [item["id"] for item in records],
                })],
                artifacts=[Artifact("automation_job", str(item["id"])) for item in records],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="list_automation_jobs", error=f"자동화 작업 목록 오류: {e}")

    def get_job_history(self, job_id: int, limit: int = 10):
        """자동화 작업 실행 기록 보기"""
        try:
            from core.scheduler import get_automation_engine
            engine = get_automation_engine()
            if engine.get_job_record(job_id) is None:
                return ToolRunResult.failed(
                    tool_name="get_job_history",
                    error=f"작업 ID {job_id}를 찾을 수 없습니다.",
                )
            records = engine.get_job_history_records(job_id, limit)
            return ToolRunResult.successful(
                tool_name="get_job_history",
                raw_output=engine.get_job_history(job_id, limit),
                evidence=[Evidence("scheduler_history", f"작업 실행 기록 {len(records)}건을 DB에서 조회했습니다.", {
                    "job_id": job_id,
                    "count": len(records),
                    "limit": limit,
                })],
                artifacts=[Artifact("automation_job", str(job_id))],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="get_job_history", error=f"실행 기록 조회 오류: {e}")

    def toggle_automation_job(self, job_id: int, enabled: bool):
        """자동화 작업 활성화/비활성화"""
        try:
            from core.scheduler import get_automation_engine
            engine = get_automation_engine()
            raw = engine.toggle_job(job_id, enabled)
            record = engine.get_job_record(job_id)
            if record is None or bool(record["enabled"]) != bool(enabled):
                return ToolRunResult.failed(
                    tool_name="toggle_automation_job",
                    error=str(raw),
                    raw_output=str(raw),
                    evidence=[Evidence("scheduler_database", "요청한 활성 상태가 DB에 반영되지 않았습니다.", {
                        "job_id": job_id,
                        "requested_enabled": bool(enabled),
                    })],
                )
            return ToolRunResult.successful(
                tool_name="toggle_automation_job",
                raw_output=str(raw),
                evidence=[Evidence("scheduler_database", "변경된 활성 상태를 DB에서 확인했습니다.", {
                    "job_id": job_id,
                    "enabled": bool(record["enabled"]),
                })],
                artifacts=[Artifact("automation_job", str(job_id))],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="toggle_automation_job", error=f"작업 상태 변경 오류: {e}")

    def delete_automation_job(self, job_id: int):
        """자동화 작업 삭제"""
        try:
            from core.scheduler import get_automation_engine
            engine = get_automation_engine()
            if engine.get_job_record(job_id) is None:
                return ToolRunResult.failed(
                    tool_name="delete_automation_job",
                    error=f"작업 ID {job_id}를 찾을 수 없습니다.",
                )
            raw = engine.delete_job(job_id)
            if engine.get_job_record(job_id) is not None:
                return ToolRunResult.failed(
                    tool_name="delete_automation_job",
                    error="작업 삭제 후에도 DB 레코드가 남아 있습니다.",
                    raw_output=str(raw),
                )
            return ToolRunResult.successful(
                tool_name="delete_automation_job",
                raw_output=str(raw),
                evidence=[Evidence("scheduler_database", "삭제 후 작업 레코드가 없음을 확인했습니다.", {
                    "job_id": job_id,
                })],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="delete_automation_job", error=f"자동화 작업 삭제 오류: {e}")

    def start_automation_engine(self):
        """자동화 엔진 시작"""
        try:
            from core.scheduler import get_automation_engine
            engine = get_automation_engine()
            raw = engine.start()
            running = engine.is_running()
            thread_alive = bool(engine.scheduler_thread and engine.scheduler_thread.is_alive())
            if not running or not thread_alive:
                return ToolRunResult.failed(
                    tool_name="start_automation_engine",
                    error=str(raw),
                    raw_output=str(raw),
                    evidence=[Evidence("automation_engine", "엔진 상태 또는 스케줄러 스레드가 실행 상태가 아닙니다.", {
                        "running": running,
                        "thread_alive": thread_alive,
                    })],
                )
            return ToolRunResult.successful(
                tool_name="start_automation_engine",
                raw_output=str(raw),
                evidence=[Evidence("automation_engine", "엔진 상태와 스케줄러 스레드 실행을 확인했습니다.", {
                    "running": True,
                    "thread_alive": True,
                })],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="start_automation_engine", error=f"자동화 엔진 시작 오류: {e}")

    def stop_automation_engine(self):
        """자동화 엔진 중지"""
        try:
            from core.scheduler import get_automation_engine
            engine = get_automation_engine()
            raw = engine.stop()
            running = engine.is_running()
            thread_alive = bool(engine.scheduler_thread and engine.scheduler_thread.is_alive())
            if running or thread_alive:
                return ToolRunResult.failed(
                    tool_name="stop_automation_engine",
                    error="자동화 엔진 또는 스케줄러 스레드가 아직 실행 중입니다.",
                    raw_output=str(raw),
                )
            return ToolRunResult.successful(
                tool_name="stop_automation_engine",
                raw_output=str(raw),
                evidence=[Evidence("automation_engine", "엔진 상태와 스케줄러 스레드 중지를 확인했습니다.", {
                    "running": False,
                    "thread_alive": False,
                })],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="stop_automation_engine", error=f"자동화 엔진 중지 오류: {e}")

    def is_automation_engine_running(self):
        """자동화 엔진 실행 여부"""
        try:
            from core.scheduler import get_automation_engine
            engine = get_automation_engine()
            running = engine.is_running()
            thread_alive = bool(engine.scheduler_thread and engine.scheduler_thread.is_alive())
            return ToolRunResult.successful(
                tool_name="is_automation_engine_running",
                raw_output="실행 중" if running else "중지됨",
                evidence=[Evidence("automation_engine", "엔진 상태와 스케줄러 스레드 상태를 조회했습니다.", {
                    "running": running,
                    "thread_alive": thread_alive,
                })],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="is_automation_engine_running", error=f"상태 확인 오류: {e}")

    # ------------------------------
    # Knowledge Graph 관련 도구 추가
    # ------------------------------
    def add_entity(self, name: str, entity_type: str, metadata: Optional[str] = None):
        """지식 그래프에 엔티티 추가"""
        try:
            from core.knowledge_graph import get_knowledge_graph, Entity
            kg = get_knowledge_graph()
            meta_dict = json.loads(metadata) if metadata else {}
            entity = Entity(name=name, entity_type=entity_type, metadata=meta_dict)
            entity_id = kg.add_entity(entity)
            saved = kg.get_entity(name, entity_type)
            if not saved or saved.id != entity_id or saved.metadata != meta_dict:
                return ToolRunResult.failed(tool_name="add_entity", error="엔티티 저장 후 재조회 검증에 실패했습니다.")
            return ToolRunResult.successful(
                tool_name="add_entity",
                raw_output=f"엔티티가 추가되었습니다: {name} ({entity_type})",
                evidence=[Evidence("knowledge_entity", "저장된 엔티티를 지식 그래프 DB에서 재조회했습니다.", {
                    "id": saved.id, "name": saved.name, "entity_type": saved.entity_type,
                    "metadata_sha256": hashlib.sha256(json.dumps(saved.metadata, sort_keys=True, ensure_ascii=False).encode()).hexdigest(),
                })],
                artifacts=[Artifact("knowledge_entity", str(saved.id), {"name": name, "entity_type": entity_type})],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="add_entity", error=f"엔티티 추가 오류: {e}")

    def get_entity(self, name: str, entity_type: Optional[str] = None):
        """지식 그래프에서 엔티티 조회"""
        try:
            from core.knowledge_graph import get_knowledge_graph
            kg = get_knowledge_graph()
            entity = kg.get_entity(name, entity_type)
            if not entity:
                return ToolRunResult.failed(tool_name="get_entity", error=f"엔티티를 찾을 수 없습니다: {name}")
            output = json.dumps(entity.to_dict(), indent=2, ensure_ascii=False)
            return ToolRunResult.successful(
                tool_name="get_entity", raw_output=output,
                evidence=[Evidence("knowledge_entity", "지식 그래프 DB에서 엔티티를 조회했습니다.", {
                    "id": entity.id, "name": entity.name, "entity_type": entity.entity_type,
                })],
                artifacts=[Artifact("knowledge_entity", str(entity.id), {"name": entity.name})],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="get_entity", error=f"엔티티 조회 오류: {e}")

    def search_entities(self, query: str, entity_type: Optional[str] = None, limit: int = 20):
        """지식 그래프에서 엔티티 검색"""
        try:
            from core.knowledge_graph import get_knowledge_graph
            kg = get_knowledge_graph()
            entities = kg.search_entities(query, entity_type, limit)
            result = [f"📋 검색 결과 ({len(entities)}건):"]
            for entity in entities:
                result.append(f"\n- {entity.name} ({entity.entity_type})")
                if entity.metadata:
                    result.append(f"  메타데이터: {json.dumps(entity.metadata, ensure_ascii=False)}")
            output = "\n".join(result) if entities else "검색 결과가 없습니다."
            return ToolRunResult.successful(
                tool_name="search_entities", raw_output=output,
                evidence=[Evidence("knowledge_query", f"지식 그래프에서 엔티티 {len(entities)}건을 조회했습니다.", {
                    "query": query, "entity_type": entity_type, "limit": limit,
                    "entity_ids": [entity.id for entity in entities],
                })],
                artifacts=[Artifact("knowledge_entity", str(entity.id), {"name": entity.name}) for entity in entities],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="search_entities", error=f"엔티티 검색 오류: {e}")

    def delete_entity(self, name: str, entity_type: Optional[str] = None):
        """지식 그래프에서 엔티티 삭제"""
        try:
            from core.knowledge_graph import get_knowledge_graph
            kg = get_knowledge_graph()
            deleted = kg.delete_entity(name, entity_type)
            if not deleted:
                return ToolRunResult.failed(tool_name="delete_entity", error=f"엔티티를 찾을 수 없습니다: {name}")
            if kg.get_entity(name, entity_type) is not None:
                return ToolRunResult.failed(tool_name="delete_entity", error="엔티티 삭제 후에도 DB 레코드가 남아 있습니다.")
            return ToolRunResult.successful(
                tool_name="delete_entity", raw_output=f"엔티티가 삭제되었습니다: {name}",
                evidence=[Evidence("knowledge_entity_absent", "삭제 후 엔티티가 조회되지 않음을 확인했습니다.", {
                    "name": name, "entity_type": entity_type,
                })],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="delete_entity", error=f"엔티티 삭제 오류: {e}")

    def add_triple(self, subject: str, predicate: str, object_: str,
                   subject_type: str = "thing", object_type: str = "thing",
                   metadata: Optional[str] = None) -> str:
        """지식 그래프에 트리플 추가"""
        try:
            from core.knowledge_graph import get_knowledge_graph
            kg = get_knowledge_graph()
            meta_dict = json.loads(metadata) if metadata else {}
            relation_id = kg.add_triple(subject, predicate, object_, subject_type, object_type, meta_dict)
            relation = next((
                item for item in kg.get_relations(subject, "from")
                if item.id == relation_id and item.relation_type == predicate and item.to_entity == object_
            ), None)
            if relation is None:
                return ToolRunResult.failed(tool_name="add_triple", error="트리플 저장 후 관계 재조회 검증에 실패했습니다.")
            return ToolRunResult.successful(
                tool_name="add_triple",
                raw_output=f"트리플이 추가되었습니다: ({subject}) -[{predicate}]-> ({object_})",
                evidence=[Evidence("knowledge_relation", "저장된 관계를 지식 그래프 DB에서 재조회했습니다.", relation.to_dict())],
                artifacts=[Artifact("knowledge_relation", str(relation.id))],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="add_triple", error=f"트리플 추가 오류: {e}")

    def get_relations(self, entity_name: str, direction: str = "both"):
        """엔티티의 관계 조회"""
        try:
            from core.knowledge_graph import get_knowledge_graph
            kg = get_knowledge_graph()
            relations = kg.get_relations(entity_name, direction)
            result = [f"🔗 {entity_name}의 관계 ({len(relations)}건):"]
            for rel in relations:
                if rel.from_entity == entity_name:
                    result.append(f"\n- → {rel.relation_type} → {rel.to_entity}")
                else:
                    result.append(f"\n- ← {rel.relation_type} ← {rel.from_entity}")
                if rel.metadata:
                    result.append(f"  메타데이터: {json.dumps(rel.metadata, ensure_ascii=False)}")
            output = "\n".join(result) if relations else f"엔티티 {entity_name}의 관계가 없습니다."
            return ToolRunResult.successful(
                tool_name="get_relations", raw_output=output,
                evidence=[Evidence("knowledge_relations", f"관계 {len(relations)}건을 DB에서 조회했습니다.", {
                    "entity_name": entity_name, "direction": direction,
                    "relation_ids": [relation.id for relation in relations],
                })],
                artifacts=[Artifact("knowledge_relation", str(relation.id)) for relation in relations],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="get_relations", error=f"관계 조회 오류: {e}")

    def get_subgraph(self, entity_name: str, depth: int = 2):
        """서브그래프 조회"""
        try:
            from core.knowledge_graph import get_knowledge_graph
            kg = get_knowledge_graph()
            subgraph = kg.get_subgraph(entity_name, depth)
            output = json.dumps(subgraph, indent=2, ensure_ascii=False)
            return ToolRunResult.successful(
                tool_name="get_subgraph", raw_output=output,
                evidence=[Evidence("knowledge_subgraph", "지식 그래프 DB에서 서브그래프를 구성했습니다.", {
                    "entity_name": entity_name, "depth": depth,
                    "entities": len(subgraph.get("entities", [])),
                    "relations": len(subgraph.get("relations", [])),
                })],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="get_subgraph", error=f"서브그래프 조회 오류: {e}")

    # ------------------------------
    # Multi-Agent 관련 도구 추가
    # ------------------------------
    def execute_multi_agent(self, query: str):
        """멀티 에이전트 파이프라인 실행 (계획 → 실행 → 반성)"""
        try:
            from core.multi_agent import get_multi_agent_orchestrator
            orchestrator = get_multi_agent_orchestrator()
            before_count = len(orchestrator.tasks)
            result = orchestrator.execute_full_pipeline(query)
            new_tasks = orchestrator.tasks[before_count:]
            failed = [task for task in new_tasks if task.status == "failed"]
            evidence = [Evidence("multi_agent_tasks", "멀티 에이전트가 생성한 작업 상태를 확인했습니다.", {
                "query_sha256": hashlib.sha256(query.encode("utf-8")).hexdigest(),
                "task_count": len(new_tasks),
                "task_ids": [task.id for task in new_tasks],
                "statuses": [task.status for task in new_tasks],
            })]
            if failed or not str(result).strip():
                return ToolRunResult.failed(
                    tool_name="execute_multi_agent",
                    error=failed[0].error if failed else "멀티 에이전트가 빈 결과를 반환했습니다.",
                    raw_output=str(result),
                    evidence=evidence,
                )
            # 하위 작업 상태만으로 최종 LLM 답변의 사실성까지 검증할 수는 없습니다.
            return ToolRunResult.unverified(
                tool_name="execute_multi_agent",
                raw_output=str(result),
                evidence=evidence,
                artifacts=[Artifact("agent_task", task.id, {"status": task.status}) for task in new_tasks],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="execute_multi_agent", error=f"멀티 에이전트 실행 오류: {e}")

    def get_task_history(self):
        """멀티 에이전트 작업 히스토리 조회"""
        try:
            from core.multi_agent import get_multi_agent_orchestrator
            orchestrator = get_multi_agent_orchestrator()
            tasks = [asdict(task) for task in orchestrator.tasks]
            return ToolRunResult.successful(
                tool_name="get_task_history",
                raw_output=json.dumps(tasks, indent=2, ensure_ascii=False),
                evidence=[Evidence("agent_task_history", f"메모리의 멀티 에이전트 작업 {len(tasks)}건을 조회했습니다.", {
                    "count": len(tasks),
                    "task_ids": [task["id"] for task in tasks],
                })],
                artifacts=[Artifact("agent_task", task["id"], {"status": task["status"]}) for task in tasks],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="get_task_history", error=f"작업 히스토리 조회 오류: {e}")

    def run_command(self, command: str):
        is_valid, error_msg, args = self.safety.validate_command(command)
        if not is_valid:
            return ToolRunResult.failed(tool_name="run_command", error=error_msg)

        try:
            # shell=True는 allowlist를 우회하는 연결 연산자 해석 위험이 있으므로 사용하지 않습니다.
            # 셸 내장 명령은 cmd.exe를 직접 허용하지 않는 한 지원하지 않습니다.
            result = subprocess.run(
                args,
                shell=False,
                capture_output=True,
                text=True,
                timeout=30,
            )
            output = []
            if result.stdout:
                output.append(f"출력:\n{result.stdout}")
            if result.stderr:
                output.append(f"에러:\n{result.stderr}")
            output.append(f"종료 코드: {result.returncode}")
            raw_output = "\n".join(output)
            execution_evidence = Evidence("process_execution", "프로세스 종료 코드와 출력을 확인했습니다.", {
                "argv": args,
                "returncode": result.returncode,
                "stdout_sha256": hashlib.sha256(result.stdout.encode("utf-8")).hexdigest(),
                "stderr_sha256": hashlib.sha256(result.stderr.encode("utf-8")).hexdigest(),
                "stdout_length": len(result.stdout),
                "stderr_length": len(result.stderr),
            })
            if result.returncode != 0:
                return ToolRunResult.failed(
                    tool_name="run_command",
                    error=f"명령이 종료 코드 {result.returncode}로 실패했습니다.",
                    raw_output=raw_output,
                    evidence=[execution_evidence],
                )
            return ToolRunResult.successful(
                tool_name="run_command",
                raw_output=raw_output,
                evidence=[execution_evidence],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="run_command", error=f"명령어 실행 오류: {e}")

    def web_search(self, query: str, num_results: int = 5):
        if DDGS is None:
            return ToolRunResult.failed(tool_name="web_search", error="DDGS 검색 라이브러리가 설치되지 않았습니다.")

        try:
            results = []
            with DDGS() as ddgs:
                for r in ddgs.text(query, max_results=num_results):
                    results.append({
                        "title": str(r.get("title", "")),
                        "url": str(r.get("href", "")),
                        "snippet": str(r.get("body", "")),
                    })
            valid_sources = [item for item in results if item["url"].startswith(("http://", "https://"))]
            if not valid_sources:
                return ToolRunResult.failed(tool_name="web_search", error="검색 결과에서 유효한 출처 URL을 찾지 못했습니다.")
            output = "\n".join(
                f"제목: {item['title']}\n링크: {item['url']}\n요약: {item['snippet']}\n{'-' * 50}"
                for item in valid_sources
            )
            return ToolRunResult.successful(
                tool_name="web_search", raw_output=output,
                evidence=[Evidence("web_sources", f"DDGS에서 유효한 출처 {len(valid_sources)}건을 조회했습니다.", {
                    "query": query, "requested_results": num_results,
                    "source_count": len(valid_sources),
                    "urls": [item["url"] for item in valid_sources],
                })],
                artifacts=[Artifact("url", item["url"], {"title": item["title"]}) for item in valid_sources],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="web_search", error=f"웹 검색 오류: {e}")

    def set_profile(self, key: str, value: str):
        try:
            self.user_profile.set(key, value)
            saved = self.user_profile.get(key)
            if saved != value:
                return ToolRunResult.failed(tool_name="set_profile", error="프로필 저장 후 재조회 값이 일치하지 않습니다.")
            return ToolRunResult.successful(
                tool_name="set_profile", raw_output=f"프로필이 저장되었습니다: {key} = {value}",
                evidence=[Evidence("profile_database", "프로필 저장 후 값을 다시 조회했습니다.", {
                    "key": key,
                    "value_sha256": hashlib.sha256(value.encode("utf-8")).hexdigest(),
                })],
                artifacts=[Artifact("profile_entry", key)],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="set_profile", error=f"프로필 저장 오류: {e}")

    def get_profile(self, key: str = ""):
        try:
            if key:
                all_profile = self.user_profile.get_all()
                if key not in all_profile:
                    return ToolRunResult.failed(tool_name="get_profile", error=f"{key}가 프로필에 없습니다.")
                value = all_profile[key]
                output = f"{key}: {value}"
                evidence_data = {"key": key, "found": True}
            else:
                all_profile = self.user_profile.get_all()
                output = self.user_profile.get_profile_summary()
                evidence_data = {"keys": sorted(all_profile), "count": len(all_profile)}
            return ToolRunResult.successful(
                tool_name="get_profile", raw_output=output,
                evidence=[Evidence("profile_database", "사용자 프로필 DB를 조회했습니다.", evidence_data)],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="get_profile", error=f"프로필 조회 오류: {e}")

    def set_preference(self, pref_key: str, pref_value: str):
        try:
            self.user_profile.set_preference(pref_key, pref_value)
            saved = self.user_profile.get_preference(pref_key)
            if saved != pref_value:
                return ToolRunResult.failed(tool_name="set_preference", error="환경설정 저장 후 재조회 값이 일치하지 않습니다.")
            return ToolRunResult.successful(
                tool_name="set_preference", raw_output=f"환경설정이 저장되었습니다: {pref_key} = {pref_value}",
                evidence=[Evidence("preference_database", "환경설정 저장 후 값을 다시 조회했습니다.", {
                    "key": pref_key,
                    "value_sha256": hashlib.sha256(pref_value.encode("utf-8")).hexdigest(),
                })],
                artifacts=[Artifact("preference_entry", pref_key)],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="set_preference", error=f"환경설정 저장 오류: {e}")

    def speak_text(self, text: str, audio_processor=None) -> str:
        print(f"[DEBUG] ToolExecutor.speak_text 호출됨: {text}")
        with self._tts_lock:
            return self._speak_text_locked(text, audio_processor)

    def _speak_text_locked(self, text: str, audio_processor=None) -> str:
        text = normalize_for_tts(text)
        print(f"[TTS] 발음 정규화: {text}")
        if self.tts_settings.selected_custom_voice:
            custom_result = self._speak_with_custom_tts(text, audio_processor)
            # 사용자가 명시적으로 고른 커스텀 음성을 다른 사람의 목소리로
            # 조용히 대체하지 않는다. 실패 원인을 그대로 알려 다시 선택할 수 있게 한다.
            return custom_result
        if self.tts_settings.selected_edge_voice:
            edge_result = self._speak_with_edge_tts(text, audio_processor)
            if not edge_result.startswith("TTS 오류:"):
                return edge_result
            print(f"[TTS] Edge 음성 실패, Windows 음성으로 대체: {edge_result}")
        if pyttsx3 is None:
            return self._speak_with_windows_speech(text, audio_processor)

        # 빈 문자열이나 공백만 있을 때 처리
        if not text or text.strip() == "":
            text = "네, 보스."

        temp_wav_path = None
        engine = None
        try:
            # 매번 새로운 TTS engine 생성 (상태 꼬임 방지)
            print("[DEBUG] pyttsx3.init() 호출 전")
            engine = pyttsx3.init()
            print("[DEBUG] pyttsx3.init() 호출 성공")

            # 한국어 음성 설정 (가능한 경우)
            voices = engine.getProperty('voices')
            print(f"[DEBUG] voices 개수: {len(voices)}")
            selected_voice = self.tts_settings.resolve_voice(voices)
            if selected_voice:
                engine.setProperty('voice', selected_voice.id)
                print(f"[DEBUG] TTS 음성 설정: {selected_voice.name}")

            # 1. TTS를 WAV 파일로 저장
            temp_wav_path = tempfile.mktemp(suffix=".wav")
            print(f"[DEBUG] TTS WAV 저장 경로: {temp_wav_path}")
            engine.save_to_file(text, temp_wav_path)
            engine.runAndWait()

            # 2. AudioProcessor로 WAV 파일 재생 + 분석
            if audio_processor:
                audio_processor.play_and_analyze_tts(temp_wav_path)
            else:
                # AudioProcessor가 없으면 그냥 재생
                engine.say(text)
                engine.runAndWait()

            print("[DEBUG] TTS 처리 완료")
            return f"음성으로 읽어주었습니다: {text}"
        except Exception as e:
            print(f"[DEBUG] TTS 오류 발생: {e}")
            if os.name == "nt":
                return self._speak_with_windows_speech(text, audio_processor)
            return f"TTS 오류: {str(e)}"
        finally:
            # engine 정리
            if engine:
                try:
                    engine.stop()
                except Exception:
                    pass

            # 임시 WAV 파일 삭제
            if temp_wav_path and os.path.exists(temp_wav_path):
                try:
                    os.unlink(temp_wav_path)
                except Exception:
                    pass

    def _speak_with_custom_tts(self, text: str, audio_processor=None) -> str:
        voice_id = self.tts_settings.selected_custom_voice
        profile = next(
            (item for item in load_custom_voice_profiles() if str(item.get("id")) == voice_id),
            None,
        )
        if profile is None:
            return f"TTS 오류: 커스텀 음성 프로필을 찾을 수 없습니다: {voice_id}"
        media_paths = []
        try:
            if audio_processor is None:
                return "TTS 오류: 커스텀 음성을 재생할 오디오 처리기가 없습니다."
            client = self._get_custom_tts_client(voice_id, profile)
            if profile.get("streaming_mode") and hasattr(audio_processor, "play_streaming_tts"):
                audio_processor.play_streaming_tts(
                    client.stream_pcm(text or f"네, {self.tts_settings.selected_address}."),
                    prebuffer_seconds=float(profile.get("prebuffer_seconds", 1.0)),
                )
                return f"음성으로 읽어드렸습니다: {text}"
            chunks = split_tts_text(
                text or f"네, {self.tts_settings.selected_address}.",
                int(profile.get("chunk_chars", 90)),
            )
            audio_queue: queue.Queue = queue.Queue(maxsize=2)
            finished = object()

            def produce_audio():
                try:
                    for chunk in chunks:
                        audio_queue.put(("audio", client.synthesize(chunk)))
                except Exception as exc:
                    audio_queue.put(("error", exc))
                finally:
                    audio_queue.put(("finished", finished))

            producer = threading.Thread(target=produce_audio, daemon=True)
            producer.start()
            while True:
                kind, value = audio_queue.get()
                if kind == "finished":
                    break
                if kind == "error":
                    raise value
                media_paths.append(value)
                # The producer synthesizes the next sentence while this one plays.
                audio_processor.play_and_analyze_tts(value)
            return f"음성으로 읽어드렸습니다: {text}"
        except Exception as exc:
            return f"TTS 오류: GPT-SoVITS 합성 실패: {exc}"
        finally:
            for media_path in media_paths:
                if media_path and os.path.exists(media_path):
                    try:
                        os.unlink(media_path)
                    except OSError:
                        pass

    def _get_custom_tts_client(self, voice_id: str, profile: dict) -> GPTSoVITSClient:
        client = self._custom_tts_clients.get(voice_id)
        if client is None:
            client = GPTSoVITSClient(profile)
            self._custom_tts_clients[voice_id] = client
        return client

    def prepare_selected_tts(self):
        """Load the selected custom voice before the first assistant response."""
        voice_id = self.tts_settings.selected_custom_voice
        if not voice_id:
            return
        profile = next(
            (item for item in load_custom_voice_profiles() if str(item.get("id")) == voice_id),
            None,
        )
        if profile is None:
            return
        try:
            self._get_custom_tts_client(voice_id, profile).ensure_running()
            print(f"[TTS] 커스텀 음성 사전 로딩 완료: {voice_id}")
        except Exception as exc:
            print(f"[TTS] 커스텀 음성 사전 로딩 실패: {voice_id}: {exc}")

    def _speak_with_edge_tts(self, text: str, audio_processor=None) -> str:
        if edge_tts is None:
            return "TTS 오류: edge-tts가 설치되지 않았습니다."
        media_path = None
        try:
            fd, media_path = tempfile.mkstemp(suffix=".mp3")
            os.close(fd)
            communicate = edge_tts.Communicate(
                text or "네, 보스.", self.tts_settings.selected_edge_voice
            )
            asyncio.run(communicate.save(media_path))
            if audio_processor is None:
                return "TTS 오류: 온라인 음성을 재생할 오디오 처리기가 없습니다."
            audio_processor.play_and_analyze_tts(media_path)
            return f"음성으로 읽어드렸습니다: {text}"
        except Exception as exc:
            return f"TTS 오류: Edge 온라인 음성 합성 실패: {exc}"
        finally:
            if media_path and os.path.exists(media_path):
                try:
                    os.unlink(media_path)
                except OSError:
                    pass

    def _speak_with_windows_speech(self, text: str, audio_processor=None) -> str:
        """pyttsx3/SAPI COM 실패 시 .NET System.Speech로 대체한다."""
        if os.name != "nt":
            return "TTS 오류: Windows System.Speech는 Windows에서만 사용할 수 있습니다."

        wav_path = None
        try:
            env = os.environ.copy()
            env["JARVIS_TTS_TEXT"] = text or "네, 보스."
            env["JARVIS_TTS_VOICE"] = self.tts_settings.selected_voice_name
            if audio_processor is not None:
                fd, wav_path = tempfile.mkstemp(suffix=".wav")
                os.close(fd)
                env["JARVIS_TTS_WAV"] = wav_path

            script = (
                "Add-Type -AssemblyName System.Speech; "
                "$s = New-Object System.Speech.Synthesis.SpeechSynthesizer; "
                "try { "
                "if ($env:JARVIS_TTS_VOICE) { $s.SelectVoice($env:JARVIS_TTS_VOICE) }; "
                "if ($env:JARVIS_TTS_WAV) { $s.SetOutputToWaveFile($env:JARVIS_TTS_WAV) }; "
                "$s.Speak($env:JARVIS_TTS_TEXT) "
                "} finally { $s.Dispose() }"
            )
            completed = subprocess.run(
                ["powershell.exe", "-NoProfile", "-NonInteractive", "-Command", script],
                capture_output=True,
                text=True,
                timeout=60,
                env=env,
                check=False,
            )
            if completed.returncode != 0:
                detail = completed.stderr.strip() or f"exit code {completed.returncode}"
                return f"TTS 오류: Windows System.Speech 실행 실패: {detail}"
            if wav_path and audio_processor is not None:
                audio_processor.play_and_analyze_tts(wav_path)
            return f"음성으로 읽어드렸습니다: {env['JARVIS_TTS_TEXT']}"
        except Exception as e:
            return f"TTS 오류: Windows System.Speech 실행 실패: {e}"
        finally:
            if wav_path and os.path.exists(wav_path):
                try:
                    os.unlink(wav_path)
                except OSError:
                    pass

    def listen(self, duration: int = 3):
        if not WHISPER_AVAILABLE:
            return ToolRunResult.failed(tool_name="listen", error="음성 인식 의존성이 설치되지 않았습니다.")

        try:
            print(f"[STT] {duration}초 동안 말씀하세요...")
            if self.hardware_manager is None:
                return ToolRunResult.failed(tool_name="listen", error="하드웨어 관리자를 초기화하지 못했습니다.")
            device_info, sample_rate = self.hardware_manager._select_microphone()
            recording = sd.rec(
                int(duration * sample_rate),
                samplerate=sample_rate,
                channels=1,
                dtype="float32",
                device=device_info["index"],
            )
            sd.wait()
            audio = self.hardware_manager._to_16khz(recording[:, 0], sample_rate)
            result = self.hardware_manager._transcribe_audio(audio)
            text = result["text"].strip()

            if text:
                return ToolRunResult.successful(
                    tool_name="listen", raw_output=f"음성 인식 결과: {text}",
                    evidence=[Evidence("speech_transcription", "입력 장치 녹음을 STT 모델로 변환했습니다.", {
                        "duration_seconds": duration,
                        "device_index": device_info["index"],
                        "device_name": device_info["name"],
                        "native_sample_rate": sample_rate,
                        "transcript_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
                    })],
                )
            return ToolRunResult.failed(
                tool_name="listen", error="음성이 인식되지 않았습니다.",
                evidence=[Evidence("speech_transcription", "녹음은 완료됐지만 전사 결과가 비어 있습니다.", {
                    "duration_seconds": duration,
                    "device_index": device_info["index"],
                    "native_sample_rate": sample_rate,
                })],
            )
        except FileNotFoundError as e:
            # FFmpeg가 설치되지 않았는지 확인
            ffmpeg_available = True
            try:
                import subprocess
                subprocess.run(["ffmpeg", "-version"], capture_output=True, timeout=5)
            except (FileNotFoundError, subprocess.TimeoutExpired):
                ffmpeg_available = False

            if not ffmpeg_available:
                return ToolRunResult.failed(tool_name="listen", error="FFmpeg가 설치되지 않았거나 PATH에 없습니다.")
            return ToolRunResult.failed(tool_name="listen", error=f"STT 오류: {e}")
        except Exception as e:
            return ToolRunResult.failed(tool_name="listen", error=f"STT 오류: {e}")

    def list_audio_input_devices(self):
        if self.hardware_manager is None:
            return ToolRunResult.failed(tool_name="list_audio_input_devices", error="마이크 기능을 초기화하지 못했습니다.")
        try:
            devices = self.hardware_manager.list_input_devices()
            default_index = sd.default.device[0] if WHISPER_AVAILABLE else None
            lines = ["사용 가능한 마이크 입력 장치:"]
            for item in devices:
                marker = " (기본)" if item["index"] == default_index else ""
                lines.append(
                    f"- {item['index']}: {item['name']} / {int(item['default_samplerate'])}Hz{marker}"
                )
            if not devices:
                lines = ["사용 가능한 마이크 입력 장치가 없습니다."]
            return ToolRunResult.successful(
                tool_name="list_audio_input_devices",
                raw_output="\n".join(lines),
                evidence=[Evidence("audio_devices", f"오디오 입력 장치 {len(devices)}개를 시스템에서 조회했습니다.", {
                    "count": len(devices),
                    "default_index": default_index,
                    "devices": [
                        {
                            "index": item["index"],
                            "name": item["name"],
                            "default_samplerate": item["default_samplerate"],
                            "max_input_channels": item.get("max_input_channels"),
                        }
                        for item in devices
                    ],
                })],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="list_audio_input_devices", error=f"마이크 장치 조회 오류: {e}")

    def list_camera_devices(self):
        if self.multimodal_manager is None:
            return ToolRunResult.failed(tool_name="list_camera_devices", error="카메라 기능을 초기화하지 못했습니다.")
        try:
            devices = self.multimodal_manager.list_camera_devices()
            output = (
                "사용 가능한 카메라:\n" + "\n".join(
                    f"- {item['index']}: {item['backend']} / {item['width']}x{item['height']}"
                    for item in devices
                )
                if devices else "프레임을 읽을 수 있는 카메라가 없습니다."
            )
            return ToolRunResult.successful(
                tool_name="list_camera_devices",
                raw_output=output,
                evidence=[Evidence("camera_devices", f"프레임 읽기 검사를 통과한 카메라 {len(devices)}개를 조회했습니다.", {
                    "count": len(devices),
                    "devices": devices,
                })],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="list_camera_devices", error=f"카메라 장치 조회 오류: {e}")

    def add_document(self, file_path: str):
        if self.rag_manager is None:
            return ToolRunResult.failed(tool_name="add_document",error="RAG 기능을 초기화할 수 없습니다.")
        is_valid, error_msg = self.safety.validate_path(file_path)
        if not is_valid:
            return ToolRunResult.failed(tool_name="add_document",error=error_msg)
        raw = self.rag_manager.add_document(file_path)
        if str(raw).startswith(("오류:", "문서 추가 오류:")):
            return ToolRunResult.failed(tool_name="add_document",error=str(raw),raw_output=str(raw))
        doc_id = os.path.basename(file_path)
        saved = self.rag_manager.documents.get(doc_id)
        if not saved:
            return ToolRunResult.failed(tool_name="add_document",error="RAG 저장소에서 추가한 문서를 확인하지 못했습니다.",raw_output=str(raw))
        return ToolRunResult.successful(
            tool_name="add_document",raw_output=str(raw),
            evidence=[Evidence("rag_document","RAG 문서 저장소에서 문서와 Chunk를 확인했습니다.",{
                "doc_id":doc_id,"source":saved.get("source"),"chunks":len(saved.get("chunks") or []),
            })],
            artifacts=[Artifact("document",str(Path(file_path).resolve()),{"rag_doc_id":doc_id})],
        )

    def search_docs(self, query: str, top_k: int = 3):
        if self.rag_manager is None:
            return ToolRunResult.failed(tool_name="search_docs",error="RAG 기능을 초기화할 수 없습니다.")
        results = self.rag_manager.search_docs(query, top_k)
        output = json.dumps(results, ensure_ascii=False, indent=2)
        return ToolRunResult.successful(
            tool_name="search_docs",raw_output=output,
            evidence=[Evidence("rag_query",f"RAG에서 근거 Chunk {len(results)}건을 조회했습니다.",{
                "query":query,"top_k":top_k,"count":len(results),
                "sources":list(dict.fromkeys(str(item.get("source","")) for item in results)),
            })],
            artifacts=[
                Artifact("document",str(item["source"]))
                for item in results if item.get("source")
            ],
        )

    def list_documents(self):
        if self.rag_manager is None:
            return ToolRunResult.failed(tool_name="list_documents",error="RAG 기능을 초기화할 수 없습니다.")
        output = self.rag_manager.list_documents()
        documents = list(self.rag_manager.documents)
        return ToolRunResult.successful(
            tool_name="list_documents",raw_output=str(output),
            evidence=[Evidence("rag_catalog",f"RAG 문서 {len(documents)}개를 조회했습니다.",{"documents":documents,"count":len(documents)})],
        )

    def add_schedule_job(self, description: str, schedule_type: str, schedule_value: str, prompt: str):
        if self.scheduler_manager is None:
            return ToolRunResult.failed(tool_name="add_schedule_job", error="스케줄러 기능을 초기화하지 못했습니다.")
        engine = self.scheduler_manager.engine
        before_ids = {item["id"] for item in engine.get_job_records()}
        raw = self.scheduler_manager.add_job(description, schedule_type, schedule_value, prompt)
        created = [item for item in engine.get_job_records() if item["id"] not in before_ids]
        if not created:
            return ToolRunResult.failed(tool_name="add_schedule_job", error=str(raw), raw_output=str(raw))
        record = created[0]
        return ToolRunResult.successful(
            tool_name="add_schedule_job", raw_output=str(raw),
            evidence=[Evidence("scheduler_database", "호환 Scheduler API 등록 결과를 DB에서 재조회했습니다.", record)],
            artifacts=[Artifact("automation_job", str(record["id"]))],
        )

    def list_schedule_jobs(self):
        if self.scheduler_manager is None:
            return ToolRunResult.failed(tool_name="list_schedule_jobs", error="스케줄러 기능을 초기화하지 못했습니다.")
        records = self.scheduler_manager.engine.get_job_records()
        return ToolRunResult.successful(
            tool_name="list_schedule_jobs", raw_output=self.scheduler_manager.list_jobs(),
            evidence=[Evidence("scheduler_database", f"호환 Scheduler API로 작업 {len(records)}건을 조회했습니다.", {
                "count": len(records), "job_ids": [item["id"] for item in records],
            })],
            artifacts=[Artifact("automation_job", str(item["id"])) for item in records],
        )

    def delete_schedule_job(self, job_id: int):
        if self.scheduler_manager is None:
            return ToolRunResult.failed(tool_name="delete_schedule_job", error="스케줄러 기능을 초기화하지 못했습니다.")
        engine = self.scheduler_manager.engine
        if engine.get_job_record(job_id) is None:
            return ToolRunResult.failed(tool_name="delete_schedule_job", error=f"작업 ID {job_id}를 찾을 수 없습니다.")
        raw = self.scheduler_manager.delete_job(job_id)
        if engine.get_job_record(job_id) is not None:
            return ToolRunResult.failed(tool_name="delete_schedule_job", error="삭제 후에도 작업이 DB에 남아 있습니다.", raw_output=str(raw))
        return ToolRunResult.successful(
            tool_name="delete_schedule_job", raw_output=str(raw),
            evidence=[Evidence("scheduler_database", "호환 Scheduler API 삭제 후 작업 부재를 확인했습니다.", {"job_id": job_id})],
        )

    def start_scheduler(self):
        if self.scheduler_manager is None:
            return ToolRunResult.failed(tool_name="start_scheduler", error="스케줄러 기능을 초기화하지 못했습니다.")
        raw = self.scheduler_manager.start_scheduler()
        engine = self.scheduler_manager.engine
        alive = bool(engine.scheduler_thread and engine.scheduler_thread.is_alive())
        if not engine.is_running() or not alive:
            return ToolRunResult.failed(tool_name="start_scheduler", error=str(raw), raw_output=str(raw))
        return ToolRunResult.successful(
            tool_name="start_scheduler", raw_output=str(raw),
            evidence=[Evidence("scheduler_thread", "호환 Scheduler 엔진과 실행 스레드를 확인했습니다.", {
                "running": True, "thread_alive": True,
            })],
        )

    def stop_scheduler(self):
        if self.scheduler_manager is None:
            return ToolRunResult.failed(tool_name="stop_scheduler", error="스케줄러 기능을 초기화하지 못했습니다.")
        raw = self.scheduler_manager.stop_scheduler()
        engine = self.scheduler_manager.engine
        alive = bool(engine.scheduler_thread and engine.scheduler_thread.is_alive())
        if engine.is_running() or alive:
            return ToolRunResult.failed(tool_name="stop_scheduler", error="Scheduler가 아직 실행 중입니다.", raw_output=str(raw))
        return ToolRunResult.successful(
            tool_name="stop_scheduler", raw_output=str(raw),
            evidence=[Evidence("scheduler_thread", "호환 Scheduler 엔진과 실행 스레드 중지를 확인했습니다.", {
                "running": False, "thread_alive": False,
            })],
        )

    def start_wakeword_detection(self):
        if self.hardware_manager is None:
            return ToolRunResult.failed(tool_name="start_wakeword_detection", error="하드웨어 기능을 초기화하지 못했습니다.")
        raw = self.hardware_manager.start_wakeword_detection()
        thread = getattr(self.hardware_manager, "wakeword_thread", None)
        running = bool(self.hardware_manager.running)
        alive = bool(thread and thread.is_alive())
        if not running or not alive:
            return ToolRunResult.failed(tool_name="start_wakeword_detection", error=str(raw), raw_output=str(raw))
        return ToolRunResult.unverified(
            tool_name="start_wakeword_detection", raw_output=str(raw),
            evidence=[Evidence("detector_thread", "감지 스레드는 시작됐지만 장치 스트림 개방은 아직 확인되지 않았습니다.", {
                "detector": "wakeword", "running": running, "thread_alive": alive,
            })],
        )

    def stop_wakeword_detection(self):
        if self.hardware_manager is None:
            return ToolRunResult.failed(tool_name="stop_wakeword_detection", error="하드웨어 기능을 초기화하지 못했습니다.")
        raw = self.hardware_manager.stop_wakeword_detection()
        thread = getattr(self.hardware_manager, "wakeword_thread", None)
        if self.hardware_manager.running or (thread and thread.is_alive()):
            return ToolRunResult.failed(tool_name="stop_wakeword_detection", error="웨이크워드 감지 스레드가 아직 실행 중입니다.", raw_output=str(raw))
        return ToolRunResult.successful(
            tool_name="stop_wakeword_detection", raw_output=str(raw),
            evidence=[Evidence("detector_thread", "웨이크워드 감지 상태와 스레드 중지를 확인했습니다.", {
                "detector": "wakeword", "running": False, "thread_alive": False,
            })],
        )

    def start_clap_detection(self):
        if self.hardware_manager is None:
            return ToolRunResult.failed(tool_name="start_clap_detection", error="하드웨어 기능을 초기화하지 못했습니다.")
        raw = self.hardware_manager.start_clap_detection()
        thread = getattr(self.hardware_manager, "clap_thread", None)
        running = bool(self.hardware_manager.running)
        alive = bool(thread and thread.is_alive())
        if not running or not alive:
            return ToolRunResult.failed(tool_name="start_clap_detection", error=str(raw), raw_output=str(raw))
        return ToolRunResult.unverified(
            tool_name="start_clap_detection", raw_output=str(raw),
            evidence=[Evidence("detector_thread", "감지 스레드는 시작됐지만 장치 스트림 개방은 아직 확인되지 않았습니다.", {
                "detector": "clap", "running": running, "thread_alive": alive,
            })],
        )

    def stop_clap_detection(self):
        if self.hardware_manager is None:
            return ToolRunResult.failed(tool_name="stop_clap_detection", error="하드웨어 기능을 초기화하지 못했습니다.")
        raw = self.hardware_manager.stop_clap_detection()
        thread = getattr(self.hardware_manager, "clap_thread", None)
        if self.hardware_manager.running or (thread and thread.is_alive()):
            return ToolRunResult.failed(tool_name="stop_clap_detection", error="박수 감지 스레드가 아직 실행 중입니다.", raw_output=str(raw))
        return ToolRunResult.successful(
            tool_name="stop_clap_detection", raw_output=str(raw),
            evidence=[Evidence("detector_thread", "박수 감지 상태와 스레드 중지를 확인했습니다.", {
                "detector": "clap", "running": False, "thread_alive": False,
            })],
        )

    def analyze_image(self, image_path: str, prompt: str = "이 이미지에 무엇이 있나요?"):
        if self.multimodal_manager is None:
            return ToolRunResult.failed(tool_name="analyze_image", error="멀티모달 기능을 초기화하지 못했습니다.")
        source = Path(image_path)
        if not source.is_file():
            return ToolRunResult.failed(tool_name="analyze_image", error=f"이미지 파일을 찾을 수 없습니다: {image_path}")
        raw = self.multimodal_manager.analyze_image(image_path, prompt)
        if str(raw).startswith(("오류:", "이미지 분석 오류:")) or not str(raw).strip():
            return ToolRunResult.failed(tool_name="analyze_image", error=str(raw), raw_output=str(raw))
        return ToolRunResult.unverified(
            tool_name="analyze_image", raw_output=str(raw),
            evidence=[Evidence("vision_input", "분석에 사용한 이미지 파일의 크기와 해시를 확인했습니다. 모델 해석 자체는 미검증입니다.", {
                "path": str(source.resolve()), "size": source.stat().st_size,
                "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                "prompt_sha256": hashlib.sha256(prompt.encode("utf-8")).hexdigest(),
            })],
            artifacts=[Artifact("image", str(source.resolve()))],
        )

    def extract_text_from_pdf(self, pdf_path: str, page_num: Optional[int] = None):
        if self.multimodal_manager is None:
            return ToolRunResult.failed(tool_name="extract_text_from_pdf", error="멀티모달 기능을 초기화하지 못했습니다.")
        source = Path(pdf_path)
        if not source.is_file():
            return ToolRunResult.failed(tool_name="extract_text_from_pdf", error=f"PDF 파일을 찾을 수 없습니다: {pdf_path}")
        raw = self.multimodal_manager.extract_text_from_pdf(pdf_path, page_num)
        if str(raw).startswith(("오류:", "PDF 분석 오류:")):
            return ToolRunResult.failed(tool_name="extract_text_from_pdf", error=str(raw), raw_output=str(raw))
        try:
            from core.multimodal import fitz
            with fitz.open(pdf_path) as document:
                total_pages = document.page_count
                selected_pages = 1 if page_num is not None else total_pages
            return ToolRunResult.successful(
                tool_name="extract_text_from_pdf", raw_output=str(raw),
                evidence=[Evidence("pdf_text_extraction", "PDF를 다시 열어 페이지 구조와 출력 해시를 확인했습니다.", {
                    "path": str(source.resolve()), "total_pages": total_pages,
                    "selected_pages": selected_pages, "page_num": page_num,
                    "output_sha256": hashlib.sha256(str(raw).encode("utf-8")).hexdigest(),
                })],
                artifacts=[Artifact("pdf", str(source.resolve()))],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="extract_text_from_pdf", error=f"PDF 추출 검증 오류: {e}", raw_output=str(raw))

    def create_excel_file(self, file_path: str, data: Optional[list] = None):
        """
        엑셀 파일을 생성합니다.
        data: 2차원 리스트 (예: [["이름", "나이"], ["철수", 30], ["영희", 25]])
        """
        try:
            # openpyxl이 설치된지 확인
            try:
                import openpyxl
            except ImportError:
                return ToolRunResult.failed(tool_name="create_excel_file", error="openpyxl 라이브러리가 설치되지 않았습니다.")

            # 새 워크북 생성
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Sheet1"

            # 데이터 입력
            if data:
                for row in data:
                    ws.append(row)

            # 파일 저장
            wb.save(file_path)
            reopened = openpyxl.load_workbook(file_path, data_only=False)
            rows = list(reopened["Sheet1"].iter_rows(values_only=True))
            expected = [tuple(row) for row in (data or [])]
            if rows != expected:
                return ToolRunResult.failed(tool_name="create_excel_file", error="저장 후 Excel 데이터 검증에 실패했습니다.")
            target = Path(file_path).resolve()
            return ToolRunResult.successful(
                tool_name="create_excel_file", raw_output=f"엑셀 파일이 생성되었습니다: {target}",
                evidence=[Evidence("excel_workbook", "저장한 통합문서를 다시 열어 시트와 셀 값을 확인했습니다.", {
                    "path": str(target), "sheet": "Sheet1", "row_count": len(rows),
                    "sha256": hashlib.sha256(target.read_bytes()).hexdigest(),
                })],
                artifacts=[Artifact("spreadsheet", str(target))],
            )

        except Exception as e:
            return ToolRunResult.failed(tool_name="create_excel_file", error=f"엑셀 파일 생성 오류: {e}")

    def write_excel_cell(self, file_path: str, sheet_name: str, cell: str, value: str):
        """
        엑셀 파일의 특정 셀에 값을 씁니다.
        """
        try:
            try:
                import openpyxl
            except ImportError:
                return ToolRunResult.failed(tool_name="write_excel_cell", error="openpyxl 라이브러리가 설치되지 않았습니다.")

            if not os.path.exists(file_path):
                return ToolRunResult.failed(tool_name="write_excel_cell", error=f"파일을 찾을 수 없습니다: {file_path}")

            wb = openpyxl.load_workbook(file_path)

            if sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
            else:
                ws = wb.create_sheet(sheet_name)

            ws[cell] = value
            wb.save(file_path)
            reopened = openpyxl.load_workbook(file_path, data_only=False)
            actual = reopened[sheet_name][cell].value
            if actual != value:
                return ToolRunResult.failed(tool_name="write_excel_cell", error="저장 후 셀 값 검증에 실패했습니다.")
            target = Path(file_path).resolve()
            return ToolRunResult.successful(
                tool_name="write_excel_cell", raw_output=f"셀 {sheet_name}!{cell}에 값을 입력했습니다.",
                evidence=[Evidence("excel_cell", "통합문서를 다시 열어 요청한 셀 값을 확인했습니다.", {
                    "path": str(target), "sheet": sheet_name, "cell": cell,
                    "value_sha256": hashlib.sha256(str(actual).encode("utf-8")).hexdigest(),
                    "file_sha256": hashlib.sha256(target.read_bytes()).hexdigest(),
                })],
                artifacts=[Artifact("spreadsheet", str(target))],
            )

        except Exception as e:
            return ToolRunResult.failed(tool_name="write_excel_cell", error=f"셀 쓰기 오류: {e}")



    def capture_camera(self, save_path: Optional[str] = None):
        """카메라에서 프레임을 캡처합니다."""
        if self.multimodal_manager is None:
            return ToolRunResult.failed(tool_name="capture_camera", error="카메라 기능을 초기화하지 못했습니다.")
        try:
            result = self.multimodal_manager.capture_camera_frame_details(save_path)
            captured = Path(result["path"])
            image_hash = hashlib.sha256(captured.read_bytes()).hexdigest()
            raw_output = (
                f"카메라 {result['camera_index']}({result['backend']}, "
                f"{result['width']}x{result['height']})의 프레임을 저장했습니다: {captured}"
            )
            return ToolRunResult.successful(
                tool_name="capture_camera",
                raw_output=raw_output,
                evidence=[Evidence("camera_capture", "저장된 카메라 이미지 파일의 크기와 해시를 확인했습니다.", {
                    **result,
                    "sha256": image_hash,
                })],
                artifacts=[Artifact("image", str(captured), {
                    "width": result["width"],
                    "height": result["height"],
                    "sha256": image_hash,
                })],
            )
        except Exception as e:
            return ToolRunResult.failed(tool_name="capture_camera", error=f"카메라 캡처 오류: {e}")

    def _request_tool_permissions(self, tool_name: str) -> tuple[bool, str]:
        """모든 진입점에서 동일하게 적용되는 중앙 권한 검사."""
        try:
            from core.permission import TOOL_PERMISSION_MAP, get_permission_manager
            permission_ids = []
            mapped = TOOL_PERMISSION_MAP.get(tool_name)
            if mapped:
                permission_ids.append(mapped)
            for schema in self.plugin_registry.get_all_tools():
                if schema.name == tool_name:
                    permission_ids.extend(schema.required_permissions)
                    break
            manager = get_permission_manager()
            for permission_id in dict.fromkeys(permission_ids):
                if not manager.request_permission(permission_id):
                    return False, f"오류: 권한이 거부되었습니다: {permission_id}"
            return True, ""
        except Exception as exc:
            return False, f"오류: 권한 확인에 실패했습니다: {exc}"

    def execute_tool(self, tool_name: str, tool_input: dict):
        started_at = time.perf_counter()
        granted, error = self._request_tool_permissions(tool_name)
        if not granted:
            return ToolRunResult.failed(
                tool_name=tool_name,
                error=error.removeprefix("오류: ").strip(),
                raw_output=error,
                duration_ms=(time.perf_counter() - started_at) * 1000,
            )
        tool_functions = {
            "read_file": self.read_file,
            "write_file": self.write_file,
            "list_directory": self.list_directory,
            "run_command": self.run_command,
            "web_search": self.web_search,
            "set_profile": self.set_profile,
            "get_profile": self.get_profile,
            "set_preference": self.set_preference,
            "speak_text": self.speak_text,
            "listen": self.listen,
            "add_document": self.add_document,
            "search_docs": self.search_docs,
            "list_documents": self.list_documents,
            "add_schedule_job": self.add_schedule_job,
            "list_schedule_jobs": self.list_schedule_jobs,
            "delete_schedule_job": self.delete_schedule_job,
            "start_scheduler": self.start_scheduler,
            "stop_scheduler": self.stop_scheduler,
            "start_wakeword_detection": self.start_wakeword_detection,
            "stop_wakeword_detection": self.stop_wakeword_detection,
            "start_clap_detection": self.start_clap_detection,
            "stop_clap_detection": self.stop_clap_detection,
            "analyze_image": self.analyze_image,
            "extract_text_from_pdf": self.extract_text_from_pdf,
            "create_excel_file": self.create_excel_file,
            "write_excel_cell": self.write_excel_cell,
            "create_directory": self.create_directory,
            "delete_directory": self.delete_directory,
            "capture_camera": self.capture_camera,
            "list_audio_input_devices": self.list_audio_input_devices,
            "list_camera_devices": self.list_camera_devices,
            # Workspace 관련 도구 추가
            "set_workspace": self.set_workspace,
            "get_workspace_info": self.get_workspace_info,
            "get_workspace_tree": self.get_workspace_tree,
            # Project Indexer 관련 도구 추가
            "index_project": self.index_project,
            "search_files": self.search_files,
            "search_symbols": self.search_symbols,
            "get_file_info": self.get_file_info,
            "get_project_tree": self.get_project_tree,
            "add_semantic_memory": self.add_semantic_memory,
            "get_semantic_memory": self.get_semantic_memory,
            "search_semantic_memory": self.search_semantic_memory,
            "delete_semantic_memory": self.delete_semantic_memory,
            "add_automation_job": self.add_automation_job,
            "list_automation_jobs": self.list_automation_jobs,
            "get_job_history": self.get_job_history,
            "toggle_automation_job": self.toggle_automation_job,
            "delete_automation_job": self.delete_automation_job,
            "start_automation_engine": self.start_automation_engine,
            "stop_automation_engine": self.stop_automation_engine,
            "is_automation_engine_running": self.is_automation_engine_running,
            "add_entity": self.add_entity,
            "get_entity": self.get_entity,
            "search_entities": self.search_entities,
            "delete_entity": self.delete_entity,
            "add_triple": self.add_triple,
            "get_relations": self.get_relations,
            "get_subgraph": self.get_subgraph,
            "execute_multi_agent": self.execute_multi_agent,
            "get_task_history": self.get_task_history,
        }

        if tool_name in tool_functions:
            try:
                raw_result = tool_functions[tool_name](**tool_input)
            except TypeError as e:
                raw_result = f"툴 파라미터 오류: {str(e)}"
        else:
            # Try plugin tools
            try:
                raw_result = self.plugin_registry.execute_tool(tool_name, tool_input)
            except Exception as e:
                raw_result = f"오류: 알 수 없는 툴 '{tool_name}' (플러그인 오류: {str(e)})"
        return self._adapt_tool_output(
            tool_name,
            tool_input,
            raw_result,
            (time.perf_counter() - started_at) * 1000,
        )

    def _adapt_tool_output(
        self,
        tool_name: str,
        tool_input: dict,
        result,
        duration_ms: float,
    ) -> ToolRunResult:
        """Plugin typed 결과는 보존하고 레거시 문자열은 공통 검증 계약으로 변환한다."""
        if isinstance(result, ToolRunResult):
            if result.duration_ms <= 0:
                result.duration_ms = max(0.0, float(duration_ms))
            return result
        verification = self.verifier.verify(tool_name, tool_input, str(result))
        return ToolRunResult.from_verification(
            tool_name=tool_name,
            raw_output=str(result),
            verification=verification,
            duration_ms=duration_ms,
        )


def get_tools_schema() -> list[dict]:
    schema = [
        {
            "name": "read_file",
            "description": "파일의 내용을 읽어옵니다",
            "input_schema": {
                "type": "object",
                "properties": {
                    "path": {
                        "type": "string",
                        "description": "읽을 파일의 경로",
                    }
                },
                "required": ["path"],
            },
        },
        {
            "name": "write_file",
            "description": "파일에 내용을 씁니다",
            "input_schema": {
                "type": "object",
                "properties": {
                    "path": {
                        "type": "string",
                        "description": "쓸 파일의 경로",
                    },
                    "content": {
                        "type": "string",
                        "description": "파일에 쓸 내용",
                    },
                },
                "required": ["path", "content"],
            },
        },
        {
            "name": "list_directory",
            "description": "디렉토리의 파일과 폴더 목록을 보여줍니다",
            "input_schema": {
                "type": "object",
                "properties": {
                    "path": {
                        "type": "string",
                        "description": "목록을 볼 디렉토리 경로",
                    }
                },
                "required": ["path"],
            },
        },
        {
            "name": "run_command",
            "description": "시스템 명령어를 실행합니다",
            "input_schema": {
                "type": "object",
                "properties": {
                    "command": {
                        "type": "string",
                        "description": "실행할 명령어",
                    }
                },
                "required": ["command"],
            },
        },
        {
            "name": "web_search",
            "description": "웹에서 정보를 검색합니다",
            "input_schema": {
                "type": "object",
                "properties": {
                    "query": {
                        "type": "string",
                        "description": "검색 쿼리",
                    },
                    "num_results": {
                        "type": "integer",
                        "description": "검색 결과 수 (기본 5개)",
                        "default": 5,
                    },
                },
                "required": ["query"],
            },
        },
        {
            "name": "set_profile",
            "description": "사용자 프로필에 정보를 저장합니다 (예: 이름, 취미, 선호도 등)",
            "input_schema": {
                "type": "object",
                "properties": {
                    "key": {
                        "type": "string",
                        "description": "프로필 키 (예: '이름', '취미', '선호_언어')",
                    },
                    "value": {
                        "type": "string",
                        "description": "프로필 값",
                    },
                },
                "required": ["key", "value"],
            },
        },
        {
            "name": "get_profile",
            "description": "사용자 프로필 정보를 조회합니다. 키를 지정하지 않으면 전체 프로필을 보여줍니다",
            "input_schema": {
                "type": "object",
                "properties": {
                    "key": {
                        "type": "string",
                        "description": "조회할 프로필 키 (선택사항)",
                    },
                },
                "required": [],
            },
        },
        {
            "name": "set_preference",
            "description": "사용자 환경설정을 저장합니다 (예: 응답_스타일, 언어 등)",
            "input_schema": {
                "type": "object",
                "properties": {
                    "pref_key": {
                        "type": "string",
                        "description": "환경설정 키",
                    },
                    "pref_value": {
                        "type": "string",
                        "description": "환경설정 값",
                    },
                },
                "required": ["pref_key", "pref_value"],
            },
        },
        {
            "name": "speak_text",
            "description": "텍스트를 음성으로 읽어줍니다 (TTS)",
            "input_schema": {
                "type": "object",
                "properties": {
                    "text": {
                        "type": "string",
                        "description": "음성으로 읽을 텍스트",
                    },
                },
                "required": ["text"],
            },
        },
        {
            "name": "listen",
            "description": "마이크에서 음성을 녹음하고 텍스트로 변환합니다 (STT)",
            "input_schema": {
                "type": "object",
                "properties": {
                    "duration": {
                        "type": "integer",
                        "description": "녹음 시간 (초, 기본 3초)",
                        "default": 3,
                    },
                },
                "required": [],
            },
        },
        {
            "name": "add_document",
            "description": "문서를 RAG 지식 베이스에 추가합니다",
            "input_schema": {
                "type": "object",
                "properties": {
                    "file_path": {
                        "type": "string",
                        "description": "추가할 문서의 파일 경로",
                    },
                },
                "required": ["file_path"],
            },
        },
        {
            "name": "search_docs",
            "description": "RAG 지식 베이스에서 관련 문서를 검색합니다",
            "input_schema": {
                "type": "object",
                "properties": {
                    "query": {
                        "type": "string",
                        "description": "검색 쿼리",
                    },
                    "top_k": {
                        "type": "integer",
                        "description": "반환할 결과 개수 (기본 3개)",
                        "default": 3,
                    },
                },
                "required": ["query"],
            },
        },
        {
            "name": "list_documents",
            "description": "RAG 지식 베이스에 저장된 문서 목록을 보여줍니다",
            "input_schema": {
                "type": "object",
                "properties": {},
                "required": [],
            },
        },
        {
            "name": "add_schedule_job",
            "description": "스케줄 작업을 추가합니다",
            "input_schema": {
                "type": "object",
                "properties": {
                    "description": {
                        "type": "string",
                        "description": "작업 설명",
                    },
                    "schedule_type": {
                        "type": "string",
                        "description": "스케줄 타입: every_minutes, every_hours, every_days, daily_at, every_weeks",
                    },
                    "schedule_value": {
                        "type": "string",
                        "description": "스케줄 값: 숫자 또는 시간(HH:MM)",
                    },
                    "prompt": {
                        "type": "string",
                        "description": "실행할 프롬프트",
                    },
                },
                "required": ["description", "schedule_type", "schedule_value", "prompt"],
            },
        },
        {
            "name": "list_schedule_jobs",
            "description": "등록된 스케줄 작업 목록을 보여줍니다",
            "input_schema": {
                "type": "object",
                "properties": {},
                "required": [],
            },
        },
        {
            "name": "delete_schedule_job",
            "description": "스케줄 작업을 삭제합니다",
            "input_schema": {
                "type": "object",
                "properties": {
                    "job_id": {
                        "type": "integer",
                        "description": "삭제할 작업 ID",
                    },
                },
                "required": ["job_id"],
            },
        },
        {
            "name": "start_scheduler",
            "description": "스케줄러를 시작합니다",
            "input_schema": {
                "type": "object",
                "properties": {},
                "required": [],
            },
        },
        {
            "name": "stop_scheduler",
            "description": "스케줄러를 중지합니다",
            "input_schema": {
                "type": "object",
                "properties": {},
                "required": [],
            },
        },
        {
            "name": "start_wakeword_detection",
            "description": "웨이크워드('자비스') 감지를 시작합니다",
            "input_schema": {
                "type": "object",
                "properties": {},
                "required": [],
            },
        },
        {
            "name": "stop_wakeword_detection",
            "description": "웨이크워드 감지를 중지합니다",
            "input_schema": {
                "type": "object",
                "properties": {},
                "required": [],
            },
        },
        {
            "name": "start_clap_detection",
            "description": "박수 감지를 시작합니다 (두 번 박수 치면 트리거)",
            "input_schema": {
                "type": "object",
                "properties": {},
                "required": [],
            },
        },
        {
            "name": "stop_clap_detection",
            "description": "박수 감지를 중지합니다",
            "input_schema": {
                "type": "object",
                "properties": {},
                "required": [],
            },
        },
        {
            "name": "analyze_image",
            "description": "이미지를 분석합니다 (메타데이터 확인)",
            "input_schema": {
                "type": "object",
                "properties": {
                    "image_path": {
                        "type": "string",
                        "description": "이미지 파일 경로",
                    },
                    "prompt": {
                        "type": "string",
                        "description": "분석 프롬프트 (선택사항)",
                        "default": "이 이미지에 무엇이 있나요?",
                    },
                },
                "required": ["image_path"],
            },
        },
        {
            "name": "extract_text_from_pdf",
            "description": "PDF에서 텍스트를 추출합니다",
            "input_schema": {
                "type": "object",
                "properties": {
                    "pdf_path": {
                        "type": "string",
                        "description": "PDF 파일 경로",
                    },
                    "page_num": {
                        "type": "integer",
                        "description": "특정 페이지 번호 (선택사항, 지정하지 않으면 전체)",
                    },
                },
                "required": ["pdf_path"],
            },
        },
        {
            "name": "create_excel_file",
            "description": "엑셀 파일을 생성합니다. data 매개변수로 2차원 리스트를 전달하면 데이터를 입력할 수 있습니다.",
            "input_schema": {
                "type": "object",
                "properties": {
                    "file_path": {
                        "type": "string",
                        "description": "생성할 엑셀 파일 경로 (예: C:/Users/ice31/Desktop/test.xlsx)",
                    },
                    "data": {
                        "type": "array",
                        "description": "2차원 리스트 데이터 (예: [['이름', '나이'], ['철수', 30]])",
                        "items": {
                            "type": "array",
                            "items": {
                                "type": "string"
                            }
                        }
                    }
                },
                "required": ["file_path"],
            },
        },
        {
            "name": "write_excel_cell",
            "description": "엑셀 파일의 특정 셀에 값을 씁니다.",
            "input_schema": {
                "type": "object",
                "properties": {
                    "file_path": {
                        "type": "string",
                        "description": "엑셀 파일 경로",
                    },
                    "sheet_name": {
                        "type": "string",
                        "description": "시트 이름",
                    },
                    "cell": {
                        "type": "string",
                        "description": "셀 위치 (예: A1, B2)",
                    },
                    "value": {
                        "type": "string",
                        "description": "입력할 값",
                    }
                },
                "required": ["file_path", "sheet_name", "cell", "value"],
            },
        },
        {
            "name": "create_directory",
            "description": "폴더를 생성합니다.",
            "input_schema": {
                "type": "object",
                "properties": {
                    "dir_path": {
                        "type": "string",
                        "description": "생성할 폴더 경로",
                    }
                },
                "required": ["dir_path"],
            },
        },
        {
            "name": "delete_directory",
            "description": "폴더를 삭제합니다 (주의: 폴더 내부 파일도 함께 삭제됨).",
            "input_schema": {
                "type": "object",
                "properties": {
                    "dir_path": {
                        "type": "string",
                        "description": "삭제할 폴더 경로",
                    }
                },
                "required": ["dir_path"],
            },
        },
        {
            "name": "capture_camera",
            "description": "카메라에서 프레임을 캡처합니다.",
            "input_schema": {
                "type": "object",
                "properties": {
                    "save_path": {
                        "type": "string",
                        "description": "저장할 경로 (선택사항, 지정하지 않으면 임시 파일에 저장)",
                    },
                },
                "required": [],
            },
        },
        {
            "name": "list_camera_devices",
            "description": "현재 실제 프레임을 읽을 수 있는 카메라 장치 목록을 확인합니다.",
            "input_schema": {"type": "object", "properties": {}, "required": []},
        },
        {
            "name": "list_audio_input_devices",
            "description": "현재 사용 가능한 마이크 입력 장치와 기본 장치를 확인합니다.",
            "input_schema": {"type": "object", "properties": {}, "required": []},
        },
        # ------------------------------
        # Workspace 관련 도구 스키마
        # ------------------------------
        {
            "name": "set_workspace",
            "description": "작업 공간을 설정합니다. 설정된 경로 내에서만 파일 작업이 가능합니다.",
            "input_schema": {
                "type": "object",
                "properties": {
                    "path": {
                        "type": "string",
                        "description": "작업 공간으로 사용할 폴더 경로",
                    }
                },
                "required": ["path"],
            },
        },
        {
            "name": "get_workspace_info",
            "description": "현재 설정된 Workspace의 정보를 반환합니다.",
            "input_schema": {
                "type": "object",
                "properties": {},
                "required": [],
            },
        },
        {
            "name": "get_workspace_tree",
            "description": "Workspace의 파일 트리 구조를 반환합니다.",
            "input_schema": {
                "type": "object",
                "properties": {},
                "required": [],
            },
        },
        # ------------------------------
        # Project Indexer 관련 도구 스키마
        # ------------------------------
        {
            "name": "index_project",
            "description": "현재 Workspace나 프로젝트 루트의 파일들을 인덱싱합니다.",
            "input_schema": {
                "type": "object",
                "properties": {},
                "required": [],
            },
        },
        {
            "name": "search_files",
            "description": "프로젝트에서 파일을 검색합니다. search_type으로 name, content, extension을 선택할 수 있습니다.",
            "input_schema": {
                "type": "object",
                "properties": {
                    "query": {
                        "type": "string",
                        "description": "검색 쿼리",
                    },
                    "search_type": {
                        "type": "string",
                        "description": "검색 타입: name, content, extension (기본 name)",
                        "default": "name",
                    },
                },
                "required": ["query"],
            },
        },
        {
            "name": "search_symbols",
            "description": "코드에서 심볼을 검색합니다 (함수, 클래스, 메서드 등).",
            "input_schema": {
                "type": "object",
                "properties": {
                    "query": {
                        "type": "string",
                        "description": "검색 쿼리",
                    },
                    "symbol_type": {
                        "type": "string",
                        "description": "심볼 타입: function, class, method, variable (선택사항)",
                    },
                },
                "required": ["query"],
            },
        },
        {
            "name": "get_file_info",
            "description": "파일의 상세 정보를 반환합니다 (크기, 수정일, 내용 미리보기, 심볼 등).",
            "input_schema": {
                "type": "object",
                "properties": {
                    "path": {
                        "type": "string",
                        "description": "파일 경로",
                    },
                },
                "required": ["path"],
            },
        },
        {
            "name": "get_project_tree",
            "description": "프로젝트의 파일 트리 구조를 반환합니다.",
            "input_schema": {
                "type": "object",
                "properties": {},
                "required": [],
            },
        },
        # ------------------------------
        # Memory 관련 도구 스키마
        # ------------------------------
        {
            "name": "add_semantic_memory",
            "description": "시맨틱 메모리를 추가합니다. 사용자 프로필, 프로젝트 정보, 기술 지식 등을 저장할 수 있습니다.",
            "input_schema": {
                "type": "object",
                "properties": {
                    "key": {
                        "type": "string",
                        "description": "메모리 키 (검색용)",
                    },
                    "content": {
                        "type": "string",
                        "description": "메모리 내용",
                    },
                    "category": {
                        "type": "string",
                        "description": "카테고리 (예: 사용자_프로필, 프로젝트_정보, 기술_지식, 기타)",
                        "default": "기타",
                    },
                },
                "required": ["key", "content"],
            },
        },
        {
            "name": "get_semantic_memory",
            "description": "키로 시맨틱 메모리를 조회합니다.",
            "input_schema": {
                "type": "object",
                "properties": {
                    "key": {
                        "type": "string",
                        "description": "메모리 키",
                    },
                },
                "required": ["key"],
            },
        },
        {
            "name": "search_semantic_memory",
            "description": "시맨틱 메모리를 검색합니다.",
            "input_schema": {
                "type": "object",
                "properties": {
                    "query": {
                        "type": "string",
                        "description": "검색 쿼리",
                    },
                    "category": {
                        "type": "string",
                        "description": "카테고리 (선택사항)",
                    },
                },
                "required": ["query"],
            },
        },
        {
            "name": "delete_semantic_memory",
            "description": "시맨틱 메모리를 삭제합니다.",
            "input_schema": {
                "type": "object",
                "properties": {
                    "key": {
                        "type": "string",
                        "description": "메모리 키",
                    },
                },
                "required": ["key"],
            },
        },
        # ------------------------------
        # Automation Engine 관련 도구 스키마
        # ------------------------------
        {
            "name": "add_automation_job",
            "description": "자동화 작업을 추가합니다. 스케줄 기반으로 LLM 작업을 자동 실행합니다.",
            "input_schema": {
                "type": "object",
                "properties": {
                    "description": {
                        "type": "string",
                        "description": "작업 설명",
                    },
                    "schedule_type": {
                        "type": "string",
                        "description": "스케줄 타입: every_minutes, every_hours, every_days, daily_at, every_weeks",
                    },
                    "schedule_value": {
                        "type": "string",
                        "description": "스케줄 값: 숫자 또는 시간(HH:MM)",
                    },
                    "prompt": {
                        "type": "string",
                        "description": "실행할 프롬프트",
                    },
                },
                "required": ["description", "schedule_type", "schedule_value", "prompt"],
            },
        },
        {
            "name": "list_automation_jobs",
            "description": "자동화 작업 목록을 보여줍니다.",
            "input_schema": {
                "type": "object",
                "properties": {},
                "required": [],
            },
        },
        {
            "name": "get_job_history",
            "description": "자동화 작업의 실행 기록을 보여줍니다.",
            "input_schema": {
                "type": "object",
                "properties": {
                    "job_id": {
                        "type": "integer",
                        "description": "작업 ID",
                    },
                    "limit": {
                        "type": "integer",
                        "description": "최근 몇 건을 보여줄지 (기본 10)",
                        "default": 10,
                    },
                },
                "required": ["job_id"],
            },
        },
        {
            "name": "toggle_automation_job",
            "description": "자동화 작업을 활성화하거나 비활성화합니다.",
            "input_schema": {
                "type": "object",
                "properties": {
                    "job_id": {
                        "type": "integer",
                        "description": "작업 ID",
                    },
                    "enabled": {
                        "type": "boolean",
                        "description": "활성화 여부 (true/false)",
                    },
                },
                "required": ["job_id", "enabled"],
            },
        },
        {
            "name": "delete_automation_job",
            "description": "자동화 작업을 삭제합니다.",
            "input_schema": {
                "type": "object",
                "properties": {
                    "job_id": {
                        "type": "integer",
                        "description": "작업 ID",
                    },
                },
                "required": ["job_id"],
            },
        },
        {
            "name": "start_automation_engine",
            "description": "자동화 엔진을 시작합니다.",
            "input_schema": {
                "type": "object",
                "properties": {},
                "required": [],
            },
        },
        {
            "name": "stop_automation_engine",
            "description": "자동화 엔진을 중지합니다.",
            "input_schema": {
                "type": "object",
                "properties": {},
                "required": [],
            },
        },
        {
            "name": "is_automation_engine_running",
            "description": "자동화 엔진이 실행 중인지 확인합니다.",
            "input_schema": {
                "type": "object",
                "properties": {},
                "required": [],
            },
        },
        # ------------------------------
        # Knowledge Graph 관련 도구 스키마
        # ------------------------------
        {
            "name": "add_entity",
            "description": "지식 그래프에 엔티티를 추가합니다.",
            "input_schema": {
                "type": "object",
                "properties": {
                    "name": {
                        "type": "string",
                        "description": "엔티티 이름",
                    },
                    "entity_type": {
                        "type": "string",
                        "description": "엔티티 타입",
                    },
                    "metadata": {
                        "type": "string",
                        "description": "메타데이터 (JSON 문자열, 선택사항)",
                    },
                },
                "required": ["name", "entity_type"],
            },
        },
        {
            "name": "get_entity",
            "description": "지식 그래프에서 엔티티를 조회합니다.",
            "input_schema": {
                "type": "object",
                "properties": {
                    "name": {
                        "type": "string",
                        "description": "엔티티 이름",
                    },
                    "entity_type": {
                        "type": "string",
                        "description": "엔티티 타입 (선택사항)",
                    },
                },
                "required": ["name"],
            },
        },
        {
            "name": "search_entities",
            "description": "지식 그래프에서 엔티티를 검색합니다.",
            "input_schema": {
                "type": "object",
                "properties": {
                    "query": {
                        "type": "string",
                        "description": "검색 쿼리",
                    },
                    "entity_type": {
                        "type": "string",
                        "description": "엔티티 타입 (선택사항)",
                    },
                    "limit": {
                        "type": "integer",
                        "description": "최대 결과 수 (기본 20)",
                        "default": 20,
                    },
                },
                "required": ["query"],
            },
        },
        {
            "name": "delete_entity",
            "description": "지식 그래프에서 엔티티를 삭제합니다.",
            "input_schema": {
                "type": "object",
                "properties": {
                    "name": {
                        "type": "string",
                        "description": "엔티티 이름",
                    },
                    "entity_type": {
                        "type": "string",
                        "description": "엔티티 타입 (선택사항)",
                    },
                },
                "required": ["name"],
            },
        },
        {
            "name": "add_triple",
            "description": "지식 그래프에 트리플(주어-술어-목적어)을 추가합니다.",
            "input_schema": {
                "type": "object",
                "properties": {
                    "subject": {
                        "type": "string",
                        "description": "주어 엔티티",
                    },
                    "predicate": {
                        "type": "string",
                        "description": "술어(관계 타입)",
                    },
                    "object_": {
                        "type": "string",
                        "description": "목적어 엔티티",
                    },
                    "subject_type": {
                        "type": "string",
                        "description": "주어 엔티티 타입 (기본: thing)",
                        "default": "thing",
                    },
                    "object_type": {
                        "type": "string",
                        "description": "목적어 엔티티 타입 (기본: thing)",
                        "default": "thing",
                    },
                    "metadata": {
                        "type": "string",
                        "description": "메타데이터 (JSON 문자열, 선택사항)",
                    },
                },
                "required": ["subject", "predicate", "object_"],
            },
        },
        {
            "name": "get_relations",
            "description": "엔티티의 관계를 조회합니다.",
            "input_schema": {
                "type": "object",
                "properties": {
                    "entity_name": {
                        "type": "string",
                        "description": "엔티티 이름",
                    },
                    "direction": {
                        "type": "string",
                        "description": "관계 방향: from, to, both (기본: both)",
                        "default": "both",
                    },
                },
                "required": ["entity_name"],
            },
        },
        {
            "name": "get_subgraph",
            "description": "엔티티를 중심으로 서브그래프를 조회합니다.",
            "input_schema": {
                "type": "object",
                "properties": {
                    "entity_name": {
                        "type": "string",
                        "description": "엔티티 이름",
                    },
                    "depth": {
                        "type": "integer",
                        "description": "탐색 깊이 (기본: 2)",
                        "default": 2,
                    },
                },
                "required": ["entity_name"],
            },
        },
        # ------------------------------
        # Multi-Agent 관련 도구 스키마
        # ------------------------------
        {
            "name": "execute_multi_agent",
            "description": "멀티 에이전트 파이프라인을 실행합니다 (계획 → 실행 → 반성).",
            "input_schema": {
                "type": "object",
                "properties": {
                    "query": {
                        "type": "string",
                        "description": "사용자 쿼리",
                    },
                },
                "required": ["query"],
            },
        },
        {
            "name": "get_task_history",
            "description": "멀티 에이전트 작업 히스토리를 조회합니다.",
            "input_schema": {
                "type": "object",
                "properties": {},
                "required": [],
            },
        },
    ]
    # Add plugin tools
    try:
        plugin_registry = get_plugin_registry()
        plugin_registry.load_plugins_from_directory()
        for tool in plugin_registry.get_all_tools():
            schema.append({
                "name": tool.name,
                "description": tool.description,
                "input_schema": tool.input_schema,
            })
    except Exception as e:
        print(f"[get_tools_schema] Plugin 스키마 로딩 오류: {e}")
    return schema


# 자율 실행 루프(Executor)나 native tool-calling(llm.py)에서 모델이 스스로
# 선택하면 안 되는 도구 목록. 단일 진실 공급원으로 여기 하나만 둡니다.
# - speak_text: main_qt.py가 최종 답변을 자동으로 TTS 재생하므로, 루프 중 모델이
#   이 도구를 스스로 호출하면 중복 발화/무한 호출 루프에 빠질 수 있습니다.
# - listen: 음성 입력을 기다리며 블로킹되므로 자율 루프 안에서 호출되면 프로그램이 멈춥니다.
AUTO_LOOP_EXCLUDED_TOOLS = ["speak_text", "listen", "execute_multi_agent", "get_task_history"]


def get_tools_description_text(exclude: Optional[list[str]] = None,
                               include: Optional[list[str]] = None) -> str:
    """
    get_tools_schema()를 사람이 읽을 수 있는 프롬프트 텍스트로 변환합니다.

    Planner / Executor / ReActAgent가 각자 하드코딩된 도구 목록을 유지하던 문제를 없애기 위한
    단일 진실 공급원(single source of truth)입니다. 새 도구를 추가하려면 get_tools_schema()에만
    등록하면 되고, 이 함수를 쓰는 모든 곳에 자동으로 반영됩니다.

    Args:
        exclude: 프롬프트에서 제외할 도구 이름 목록 (예: UI 전용 도구 등)
    """
    exclude_set = set(exclude or [])
    include_set = set(include) if include is not None else None
    lines = []
    for tool in get_tools_schema():
        if tool["name"] in exclude_set:
            continue
        if include_set is not None and tool["name"] not in include_set:
            continue
        # 파라미터 이름까지 같이 보여줘야 LLM이 tool_input을 정확히 채울 수 있음
        properties = tool.get("input_schema", {}).get("properties", {})
        required = set(tool.get("input_schema", {}).get("required", []))
        if properties:
            params = ", ".join(
                f"{name}{'*' if name in required else ''}"
                for name in properties.keys()
            )
            lines.append(f"- {tool['name']}({params}): {tool['description']}")
        else:
            lines.append(f"- {tool['name']}(): {tool['description']}")
    return "\n".join(lines)


def get_tool_names(exclude: Optional[list[str]] = None) -> list[str]:
    """등록된 모든 도구 이름 목록 (exclude로 일부 제외 가능)"""
    exclude_set = set(exclude or [])
    return [t["name"] for t in get_tools_schema() if t["name"] not in exclude_set]


_tools_executor = None


def get_tool_executor() -> ToolExecutor:
    global _tools_executor
    if _tools_executor is None:
        _tools_executor = ToolExecutor()
    return _tools_executor
