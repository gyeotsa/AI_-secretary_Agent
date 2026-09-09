"""Artifact-level QA for actual specialist completion, not tool success labels."""
import json
import sys

from docx import Document
from openpyxl import Workbook
import pytest

from core.coding_agent import CodingAgent, FileEdit
from core.executor import ExecutionOutcome
from core.specialist_acceptance import inspect_local_artifact, review_requirement_fulfillment, _diff_matches_content
from core.specialist_team import SpecialistTeamRuntime
from core.specialist_workspaces import get_specialist_workspace_registry
from core.tool_result import Artifact, Evidence, ToolRunResult


def contract(key):
    return get_specialist_workspace_registry().get(key).execution_contract()


def payload(key, raw=None, *, artifacts=(), evidence_kind="actual_result"):
    return {"status": "completed", "tool_status": "succeeded", "tool_name": key,
            "evidence": [{"kind": evidence_kind, "summary": "verified tool response", "data": {}}],
            "artifacts": list(artifacts),
            "tool_outputs": [{"tool_name": key, "status": "succeeded",
                              "raw_output": json.dumps(raw, ensure_ascii=False) if raw is not None else ""}]}


def test_research_url_open_does_not_pass_report_or_claim_criteria():
    value = payload("browser_open_url", artifacts=[{"kind": "url", "uri": "https://example.org"}])
    review = SpecialistTeamRuntime._review_execution(value, contract("research"))
    assert not review["passed"]
    assert all(not criterion["verified"] for criterion in review["criteria_results"])
    assert not review["readiness_checks"]["verified_evidence"]


def research_payload():
    return payload("browser_research", {
        "query": "새 제품 출시일", "sources": [{"source_id": "S1", "url": "https://example.org",
        "title": "발표문", "content": "새 제품은 2026년 9월 8일 출시됩니다.", "status_code": 200}],
        "claims": [{"text": "새 제품은 2026년 9월 8일 출시됩니다.", "citations": ["S1"]}],
    }, artifacts=[{"kind": "url", "uri": "https://example.org"}])


def test_research_real_report_requires_body_and_grounded_claim_links():
    assert SpecialistTeamRuntime._review_execution(research_payload(), contract("research"))["passed"]


@pytest.mark.parametrize("mutation", ["unknown_source", "invented_claim", "empty_body", "mismatched_hash", "no_claims"])
def test_research_rejects_content_or_citation_failures(mutation):
    value = research_payload()
    report = json.loads(value["tool_outputs"][0]["raw_output"])
    if mutation == "unknown_source":
        report["claims"][0]["citations"] = ["S999"]
    elif mutation == "invented_claim":
        report["claims"][0]["text"] = "이 제품은 세계 최고입니다."
    elif mutation == "empty_body":
        report["sources"][0]["content"] = ""
    elif mutation == "mismatched_hash":
        report["sources"][0]["content_sha256"] = "a" * 64
    else:
        report["claims"] = []
    value["tool_outputs"][0]["raw_output"] = json.dumps(report, ensure_ascii=False)
    assert not SpecialistTeamRuntime._review_execution(value, contract("research"))["passed"]


def test_document_reopens_real_docx_and_detects_content_mismatch(tmp_path):
    target = tmp_path / "report.docx"
    doc = Document(); doc.add_paragraph("검증된 실제 보고서 본문"); doc.save(target)
    value = payload("word_create_document", artifacts=[{"kind": "document", "uri": str(target)}])
    specification = contract("document")
    specification["content_requirements"] = [{"kind": "contains", "value": "검증된 실제 보고서 본문"}]
    checked = SpecialistTeamRuntime._review_execution(value, specification)
    assert checked["passed"]
    assert checked["criteria_results"][0]["inspected_artifacts"][0]["sha256"]
    doc = Document(); doc.add_paragraph("요청하지 않은 다른 내용"); doc.save(target)
    assert not SpecialistTeamRuntime._review_execution(value, specification)["passed"]


def test_document_empty_xlsx_is_not_a_completed_spreadsheet(tmp_path):
    target = tmp_path / "empty.xlsx"
    wb = Workbook(); wb.save(target)
    value = payload("excel_create_workbook", artifacts=[{"kind": "spreadsheet", "uri": str(target)}])
    review = SpecialistTeamRuntime._review_execution(value, contract("document"))
    assert not review["passed"]
    assert review["criteria_results"][1]["verified"]
    assert not review["criteria_results"][0]["verified"]
    wb.active["A1"] = "실제 데이터"; wb.save(target)
    assert SpecialistTeamRuntime._review_execution(value, contract("document"))["passed"]


def test_coding_direct_runtime_patch_and_static_command_are_reviewed(tmp_path):
    target = tmp_path / "sample.py"
    target.write_text("value = 1\n", encoding="utf-8")
    result = CodingAgent(tmp_path).apply_transaction([FileEdit("sample.py", "value = 1", "value = 2")],
        [[sys.executable, "-m", "py_compile", str(target)]])
    assert result.succeeded
    value = payload("coding_apply_patch", {
        "changed_files": result.changed_files, "diff": result.diff, "validation_output": result.validation_output,
    }, artifacts=[{"kind": "file", "uri": str(target), "metadata": {"changed": True}}])
    review = SpecialistTeamRuntime._review_execution(value, contract("coding"))
    assert review["passed"]
    assert review["criteria_results"][1]["validation_records"][0]["returncode"] == 0
    target.write_text("value = 999\n", encoding="utf-8")
    assert not SpecialistTeamRuntime._review_execution(value, contract("coding"))["passed"]
    # Reopen the current file; an old successful command cannot hide broken code.
    target.write_text("value = [\n", encoding="utf-8")
    assert not SpecialistTeamRuntime._review_execution(value, contract("coding"))["passed"]


def test_coding_file_exists_without_diff_or_tests_is_not_complete(tmp_path):
    target = tmp_path / "a.py"; target.write_text("value = 1\n", encoding="utf-8")
    value = payload("filesystem_write_file", artifacts=[{"kind": "file", "uri": str(target)}])
    assert not SpecialistTeamRuntime._review_execution(value, contract("coding"))["passed"]


class ReviewClient:
    def __init__(self, *, passed=True, invented_quote=False, missing=False):
        self.passed, self.invented_quote, self.missing = passed, invented_quote, missing
        self.calls = 0

    def chat_structured(self, messages, json_schema=None):
        self.calls += 1
        request = json.loads(messages[-1]["content"])
        if self.missing:
            return {"criteria": []}
        artifact = request["artifacts"][0]
        return {"criteria": [{"index": item["index"], "passed": self.passed,
            "reason": "본문을 확인했습니다." if self.passed else "요청한 내용을 충족하지 않습니다.",
            "evidence": [{"artifact_id": artifact["id"],
                          "quote": "존재하지 않는 인용" if self.invented_quote else artifact["content"][:100]}]}
            for item in request["criteria"]]}


@pytest.mark.parametrize("options", [{"invented_quote": True}, {"missing": True}, {"passed": False}])
def test_semantic_review_does_not_accept_ungrounded_or_missing_verdict(tmp_path, options):
    target = tmp_path / "report.txt"; target.write_text("요청에 맞는 보고서 내용입니다.", encoding="utf-8")
    value = payload("filesystem_write_file", artifacts=[{"kind": "file", "uri": str(target)}])
    client = ReviewClient(**options)
    verdict = review_requirement_fulfillment(value, contract("document"),
        {"acceptance": ["보고서 내용 검수"]}, "보고서 작성", client_provider=lambda: client)
    assert not verdict["passed"] and verdict["status"] == "needs_review"
    assert client.calls == 1


def test_semantic_review_proves_quotes_against_actual_artifact_not_tool_prose(tmp_path):
    target = tmp_path / "report.txt"; target.write_text("검증 대상 본문", encoding="utf-8")
    client = ReviewClient()
    verdict = review_requirement_fulfillment(payload("write", artifacts=[{"kind": "file", "uri": str(target)}]),
        contract("document"), {"acceptance": ["본문 포함"]}, "검증 대상 본문을 기록해", client_provider=lambda: client)
    assert verdict["passed"]
    assert all(row["grounded"] for row in verdict["criteria_results"])


def team_fixture(monkeypatch, client):
    names = ("word_create_document", "excel_create_workbook", "hwpx_create_document",
             "powerpoint_create_presentation", "pdf_create_document")
    team = SpecialistTeamRuntime(tool_name_provider=lambda: names, review_client_provider=lambda: client)
    monkeypatch.setattr(team, "_build_specialist_plan", lambda *args: {
        "planning_status": "ready", "goal": "문서 작성", "acceptance": ["본문 충족"]})
    return team


@pytest.mark.parametrize("semantic_pass", [True, False])
def test_actual_team_runtime_propagates_semantic_failure_to_returned_outcome(tmp_path, monkeypatch, semantic_pass):
    target = tmp_path / "out.docx"
    doc = Document(); doc.add_paragraph("사용자가 요청한 검증 보고서"); doc.save(target)
    original = ExecutionOutcome("문서 완료", status="completed", next_goal="후속 작업")
    original.tool_result = ToolRunResult.successful(tool_name="word_create_document", raw_output=str(target),
        evidence=[Evidence("docx_structure", "실제 저장 문서를 재열었습니다.", {"paragraphs": 1})],
        artifacts=[Artifact("document", str(target))])
    team = team_fixture(monkeypatch, ReviewClient(passed=semantic_pass))
    result, run = team.execute_workspace_request("document", "검증 보고서를 작성해줘",
        invoke_executor=lambda _: original, release_models=False)
    assert original.status == "completed"  # Original execution receipt is immutable from the caller's perspective.
    assert result.status == ("completed" if semantic_pass else "partial")
    assert run.artifacts["quality_verdict"]["passed"] is semantic_pass
    if not semantic_pass:
        assert result.next_goal == ""
        assert "완료로 처리하지 않았습니다" in result.response
        assert run.status == "degraded"
        assert run.events[-1]["status"] == "degraded"


@pytest.mark.parametrize("status", ["awaiting_user", "awaiting_input", "awaiting_approval", "cancelled"])
def test_nonterminal_or_cancelled_outcome_does_not_invoke_semantic_review(monkeypatch, status):
    client = ReviewClient()
    team = team_fixture(monkeypatch, client)
    original = ExecutionOutcome("추가 정보가 필요합니다", status=status)
    result, run = team.execute_workspace_request("document", "보고서 작성",
        invoke_executor=lambda _: original, release_models=False)
    assert result is original and result.status == status
    assert client.calls == 0
    assert run.status == status


@pytest.mark.parametrize("status_code", [None, True, -1, 100, 304, 404])
def test_research_rejects_missing_or_non_success_http_status(status_code):
    value = research_payload()
    report = json.loads(value["tool_outputs"][0]["raw_output"])
    report["sources"][0]["status_code"] = status_code
    value["tool_outputs"][0]["raw_output"] = json.dumps(report)
    assert not SpecialistTeamRuntime._review_execution(value, contract("research"))["passed"]


def test_document_does_not_hide_missing_second_deliverable(tmp_path):
    target = tmp_path / "good.txt"; target.write_text("검증 대상 보고서", encoding="utf-8")
    value = payload("write", artifacts=[{"kind": "file", "uri": str(target)},
                                      {"kind": "file", "uri": str(tmp_path / "missing.txt")}])
    assert not SpecialistTeamRuntime._review_execution(value, contract("document"))["passed"]


def test_semantic_review_rejects_file_changed_during_model_call(tmp_path):
    target = tmp_path / "report.txt"; target.write_text("검증 대상 본문", encoding="utf-8")

    class MutatingClient(ReviewClient):
        def chat_structured(self, messages, json_schema=None):
            response = super().chat_structured(messages, json_schema)
            target.write_text("사용자에 의해 변경된 본문", encoding="utf-8")
            return response

    verdict = review_requirement_fulfillment(payload("write", artifacts=[{"kind": "file", "uri": str(target)}]),
        contract("document"), {"acceptance": ["본문 포함"]}, "검증 대상 본문을 기록해", client_provider=MutatingClient)
    assert not verdict["passed"]
    assert "검수 중 산출물이 변경" in verdict["reason"]


def test_semantic_document_review_uses_text_not_preview_image_or_container(tmp_path):
    from PIL import Image
    target = tmp_path / "report.txt"; target.write_text("검증 대상 본문", encoding="utf-8")
    preview = tmp_path / "preview.png"; Image.new("RGB", (2, 2)).save(preview)
    verdict = review_requirement_fulfillment(payload("write", artifacts=[
        {"kind": "directory", "uri": str(tmp_path)}, {"kind": "file", "uri": str(target)},
        {"kind": "image", "uri": str(preview)}]), contract("document"),
        {"acceptance": ["본문 포함"]}, "검증 대상 본문을 기록해", client_provider=ReviewClient)
    assert verdict["passed"]
    assert verdict["nontext_artifacts_not_semantically_reviewed"] == [str(preview)]


@pytest.mark.parametrize("acceptance", ["wrong type", [{"criterion": "invalid"}], ["x"] * 9 + [str(i) for i in range(9)]])
def test_semantic_invalid_or_excessive_criteria_do_not_call_model(acceptance):
    client = ReviewClient()
    verdict = review_requirement_fulfillment({}, contract("document"), {"acceptance": acceptance},
        "검증", client_provider=lambda: client)
    assert not verdict["passed"] and verdict["attempts"] == 0
    assert client.calls == 0


@pytest.mark.parametrize("content,expected", [("value = 2\n", True), ("value = 1\n", False)])
def test_diff_new_side_is_bound_to_actual_file_content(content, expected):
    diff = "--- a/a.py\n+++ b/a.py\n@@ -1 +1 @@\n-value = 1\n+value = 2\n"
    assert _diff_matches_content(diff, "a.py", content) is expected


def test_coding_rejects_malformed_validation_record_alongside_success(tmp_path):
    target = tmp_path / "a.py"; target.write_text("value = 2\n", encoding="utf-8")
    value = payload("coding_apply_patch", {"changed_files": ["a.py"],
        "diff": "--- a/a.py\n+++ b/a.py\n@@ -1 +1 @@\n-value = 1\n+value = 2\n",
        "validation_output": json.dumps({"command": ["python", "-m", "py_compile", "a.py"], "returncode": 0}) + "\nFAILED MALFORMED RECORD"},
        artifacts=[{"kind": "file", "uri": str(target), "metadata": {"changed": True}}])
    assert not SpecialistTeamRuntime._review_execution(value, contract("coding"))["passed"]


@pytest.mark.parametrize("missing", ["inputs", "constraints", "acceptance", "tools"])
def test_plan_normalization_does_not_invent_required_fields(missing):
    plan = {"goal": "보고서 작성", "inputs": [], "constraints": [],
            "acceptance": ["본문 검수"], "tools": []}
    del plan[missing]
    normalized = SpecialistTeamRuntime._normalize_specialist_plan(plan)
    assert missing not in normalized
    assert not SpecialistTeamRuntime._valid_specialist_plan(normalized)


@pytest.mark.parametrize("invalid", [None, "", 7, {}, {"unsupported": "value"}])
def test_plan_normalization_preserves_invalid_entries_for_rejection(invalid):
    plan = {"goal": "보고서 작성", "inputs": [invalid], "constraints": [],
            "acceptance": ["본문 검수"], "tools": []}
    assert not SpecialistTeamRuntime._valid_specialist_plan(
        SpecialistTeamRuntime._normalize_specialist_plan(plan))


@pytest.mark.parametrize("goal", [None, 7, {"goal": "보고서 작성"}, ["보고서 작성"]])
def test_plan_requires_a_real_string_goal(goal):
    plan = {"goal": goal, "inputs": [], "constraints": [],
            "acceptance": ["본문 검수"], "tools": []}
    assert not SpecialistTeamRuntime._valid_specialist_plan(
        SpecialistTeamRuntime._normalize_specialist_plan(plan))


def isolated_supervisor(tmp_path, monkeypatch):
    from contextlib import nullcontext
    from core.task_contracts import SupervisorRuntime, TaskContractStore
    supervisor = SupervisorRuntime(TaskContractStore(str(tmp_path / "contracts.db")))
    # The persistent lifecycle is real; no model or GPU is needed for a
    # deterministic reviewer return value in this contract-focused test.
    monkeypatch.setattr(supervisor, "resource_guard", lambda _contract: nullcontext())
    return supervisor


@pytest.mark.parametrize("passed,expected", [(True, "completed"), (False, "failed")])
def test_reviewer_persistent_contract_records_actual_acceptance(tmp_path, monkeypatch, passed, expected):
    supervisor = isolated_supervisor(tmp_path, monkeypatch)
    team = SpecialistTeamRuntime(supervisor=supervisor)
    verdict = {"passed": passed, "status": "completed" if passed else "needs_review",
               "reason": "요구사항 대조 결과"}
    with team.run("document", "보고서를 작성해줘", artifacts={"document": "execution receipt"}) as run:
        assert team.execute_role(run, "reviewer", lambda _run: verdict) is verdict
    stored = supervisor.store.list_recent()[0]
    assert stored.status.value == expected
    assert stored.attempts == 1
    assert stored.failure_reason == ("" if passed else verdict["reason"])
    assert stored.evidence[-1]["kind"] == "quality_acceptance"
    assert stored.evidence[-1]["passed"] is passed
    assert stored.artifacts[0]["kind"] == "quality_verdict"


def test_non_reviewer_role_does_not_treat_arbitrary_passed_field_as_acceptance(tmp_path, monkeypatch):
    supervisor = isolated_supervisor(tmp_path, monkeypatch)
    team = SpecialistTeamRuntime(supervisor=supervisor)
    with team.run("document", "보고서를 작성해줘", artifacts={"requirements": "request"}) as run:
        team.execute_role(run, "author", lambda _run: {"passed": False, "status": "metadata"})
    assert supervisor.store.list_recent()[0].status.value == "completed"


def test_reviewer_returned_cancellation_remains_cancelled_in_persistent_contract(tmp_path, monkeypatch):
    supervisor = isolated_supervisor(tmp_path, monkeypatch)
    team = SpecialistTeamRuntime(supervisor=supervisor)
    with team.run("document", "보고서를 작성해줘", artifacts={"document": "execution receipt"}) as run:
        team.execute_role(run, "reviewer", lambda _run: {
            "passed": False, "status": "cancelled", "reason": "사용자가 취소했습니다."})
    assert supervisor.store.list_recent()[0].status.value == "cancelled"
    assert run.status == "cancelled"
    assert run.events[-1]["status"] == "cancelled"


def test_reviewer_inflight_cancellation_cannot_be_overwritten_by_negative_verdict(tmp_path, monkeypatch):
    supervisor = isolated_supervisor(tmp_path, monkeypatch)
    team = SpecialistTeamRuntime(supervisor=supervisor)

    def cancel_during_review(_run):
        current = supervisor.store.list_recent()[0]
        supervisor.request_cancel(current.contract_id)
        return {"passed": False, "status": "needs_review", "reason": "요구사항 불일치"}

    with team.run("document", "보고서를 작성해줘", artifacts={"document": "execution receipt"}) as run:
        with pytest.raises(RuntimeError, match="취소"):
            team.execute_role(run, "reviewer", cancel_during_review)
    assert supervisor.store.list_recent()[0].status.value == "cancelled"
    assert run.status == "cancelled"
    assert run.events[-1]["status"] == "cancelled"


@pytest.mark.integration
def test_live_local_semantic_acceptance_positive_and_negative(tmp_path):
    """Opt-in: two bounded local-model calls; no executor, accounts, or sends."""
    from urllib.parse import urlparse
    from core.llm import OllamaClient

    client = OllamaClient("reasoning")
    assert urlparse(client.base_url).hostname in {"localhost", "127.0.0.1", "::1"}
    target = tmp_path / "acceptance_report.txt"
    specification = contract("document")
    plan = {"acceptance": ["회의 시간이 오후 3시라고 명시되어 있다"]}
    request = "보고서에 회의 시간이 오후 3시라고 적어줘."
    results = []
    try:
        for content, expected in [("회의 시간은 오후 3시입니다.", True), ("회의 시간은 오전 9시입니다.", False)]:
            target.write_text(content, encoding="utf-8")
            value = payload("write", artifacts=[{"kind": "file", "uri": str(target)}])
            deterministic = SpecialistTeamRuntime._review_execution(value, specification)
            assert deterministic["passed"]  # Both are real, nonempty files; only one meets the request.
            review = review_requirement_fulfillment(value, specification, plan, request,
                                                    client_provider=lambda: client)
            results.append({"model": client.model, "content": content, "expected": expected,
                            "review": review})
            assert review["attempts"] == 1
            assert review["passed"] is expected, json.dumps(review, ensure_ascii=False)
    finally:
        print(json.dumps({"local_acceptance_results": results}, ensure_ascii=True))
        client.release()
